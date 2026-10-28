"""Bounded XLS/XLSX import; TS BOM LIST has a strict, isolated import path."""
import io
import re
import zipfile
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
import xlrd
from pydantic import ValidationError

from .models import BOMItem, millimeters
from .materials import compact

MAX_UPLOAD = 10 * 1024 * 1024
MAX_ROWS = 10000
FIELDS = ('material', 'spec', 'length', 'quantity')
ALIASES = {
    'material': ['재질', '재료', '강종', '材質', 'material', 'grade', 'steelgrade', 'matl'],
    'spec': ['규격', '단면', '형강규격', '부재규격', 'spec', 'specification', 'section', 'size', 'profile', 'description'],
    'length': ['길이', '길이mm', '절단길이', '절단길이mm', '부재길이', 'length', 'lengthmm', 'cutlength'],
    'quantity': ['수량', '수량ea', '개수', '본수', '소요수량', 'qty', 'quantity', 'count', 'pcs'],
}
METADATA_ALIASES = {'item_class': ['item'], 'title': ['title'], 'weight': ['weightkg'], 'paint_area': ['paintm']}
TOTAL_LABELS = {'합계', '총합계', '총계', '소계', 'total', 'grandtotal', 'subtotal'}


def normalized(value):
    return re.sub(r'[^a-z0-9가-힣材質]', '', str(value or '').lower())


def workbook(data):
    if len(data) > MAX_UPLOAD:
        raise ValueError('엑셀 파일은 10MB 이하만 업로드할 수 있습니다.')
    try:
        if data.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'):
            return XlsWorkbook(xlrd.open_workbook(file_contents=data, on_demand=True))
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError('압축 해제 크기가 너무 큰 엑셀 파일입니다.')
        return load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, xlrd.XLRDError) as exc:
        raise ValueError('올바른 .xls 또는 .xlsx 파일을 업로드해 주세요.') from exc


class XlsWorksheet:
    def __init__(self, sheet):
        self.sheet = sheet
        self.title = sheet.name
        self.max_row = sheet.nrows

    def iter_rows(self, min_row=1, max_row=None, max_col=100, values_only=True):
        for row in range(min_row - 1, min(max_row or self.max_row, self.max_row)):
            cells = self.sheet.row(row)
            yield tuple(None if col >= len(cells) or cells[col].ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK)
                        else bool(cells[col].value) if cells[col].ctype == xlrd.XL_CELL_BOOLEAN
                        else xlrd.error_text_from_code.get(cells[col].value, '#ERROR!') if cells[col].ctype == xlrd.XL_CELL_ERROR
                        else cells[col].value for col in range(max_col))


class XlsWorkbook:
    def __init__(self, book):
        self.book = book
        self.sheetnames = book.sheet_names()

    def __getitem__(self, name):
        return XlsWorksheet(self.book.sheet_by_name(name))

    def close(self):
        self.book.release_resources()


def strict_ts_bom(wb):
    if 'TS BOM LIST' in wb.sheetnames:
        return True
    if isinstance(wb, XlsWorkbook):
        raise ValueError('.xls 파일에는 TS BOM LIST 시트가 반드시 있어야 합니다. 다른 시트로 대체하지 않습니다.')
    return False


def find_mapping(headers):
    result = {}
    for field, names in ALIASES.items():
        matches = [i for i, cell in enumerate(headers) if normalized(cell) in names]
        if len(matches) == 1:
            result[field] = matches[0]
    return result


def inspect_bom(data):
    wb = workbook(data)
    try:
        if strict_ts_bom(wb):
            from .ts_bom import inspect_ts_bom
            return inspect_ts_bom(wb)
        sheets = []
        for ws in wb.worksheets[:30]:
            best = (0, 1, [], {})
            for index, row in enumerate(ws.iter_rows(max_row=30, max_col=100, values_only=True), 1):
                headers = [str(c).strip() if c is not None else '' for c in row]
                mapping = find_mapping(headers)
                if len(mapping) > best[0]:
                    best = (len(mapping), index, headers, mapping)
            _, header_row, headers, mapping = best
            sheets.append({'name': ws.title, 'header_row': header_row, 'headers': headers, 'mapping': mapping})
        if not sheets:
            raise ValueError('읽을 수 있는 시트가 없습니다.')
        suggested = next((s for s in sheets if compact(s['name']).upper() == 'TSASSEMLIST'),
                         max(sheets, key=lambda s: len(s['mapping'])))
        return {'sheets': sheets, 'suggested_sheet': suggested['name']}
    finally:
        wb.close()


def read_headers(data, sheet, header_row):
    if not 1 <= header_row <= 100:
        raise ValueError('헤더 행은 1~100 사이로 설정해 주세요.')
    wb = workbook(data)
    try:
        if strict_ts_bom(wb):
            from .ts_bom import locked_headers
            return locked_headers(wb, sheet, header_row)
        if sheet not in wb.sheetnames:
            raise ValueError('선택한 시트가 없습니다.')
        headers = next(wb[sheet].iter_rows(min_row=header_row, max_row=header_row, max_col=100, values_only=True))
        headers = [str(c).strip() if c is not None else '' for c in headers]
        return {'headers': headers, 'mapping': find_mapping(headers)}
    finally:
        wb.close()


def parse_bom(data, sheet, header_row, mapping):
    if set(mapping) != set(FIELDS) or len(set(mapping.values())) != 4:
        raise ValueError('재질·규격·길이·수량에 서로 다른 열을 지정해 주세요.')
    if any(type(v) is not int or not 0 <= v < 100 for v in mapping.values()):
        raise ValueError('열 번호가 올바르지 않습니다.')
    if not 1 <= header_row <= 100:
        raise ValueError('헤더 행은 1~100 사이로 설정해 주세요.')
    wb = workbook(data)
    try:
        if strict_ts_bom(wb):
            from .ts_bom import parse_ts_bom
            return parse_ts_bom(wb, data, sheet, header_row, mapping)
        if sheet not in wb.sheetnames:
            raise ValueError('선택한 시트가 없습니다.')
        ws = wb[sheet]
        if ws.max_row and ws.max_row > header_row + MAX_ROWS:
            raise ValueError('시트의 사용 범위가 10,000행을 초과합니다. 불필요한 행을 삭제해 주세요. 일부 데이터만 계산하지 않습니다.')
        header = next(ws.iter_rows(min_row=header_row, max_row=header_row, max_col=100, values_only=True))
        metadata = {name: matches[0] for name, aliases in METADATA_ALIASES.items()
                    if len(matches := [i for i, cell in enumerate(header) if normalized(cell) in aliases]) == 1}
        # Locate a labelled trailing total without hiding a part whose TITLE happens to be "TOTAL".
        raw_rows = list(ws.iter_rows(min_row=header_row + 1, max_row=header_row + MAX_ROWS + 1,
                                     max_col=max([*mapping.values(), *metadata.values()]) + 1, values_only=True))
        populated = [i for i, row in enumerate(raw_rows) if any(c is not None and str(c).strip() for c in row)]
        last_data_index = populated[-1] if populated else -1
        items, errors, skipped_rows = [], [], []
        for index, row in enumerate(raw_rows):
            rownum = header_row + 1 + index
            values = {key: row[col] for key, col in mapping.items()}
            if all(v is None or str(v).strip() == '' for v in row):
                continue
            if rownum > header_row + MAX_ROWS:
                raise ValueError('데이터 행은 최대 10,000행까지 읽을 수 있습니다.')
            labels = [row[c] for c in set([mapping['material'], mapping['spec'], *metadata.values()])]
            if index == last_data_index and any(normalized(v) in TOTAL_LABELS for v in labels):
                # Require no real section, or no cutting length: a real part named TOTAL remains data.
                if not compact(values['spec']) or normalized(values['spec']) in TOTAL_LABELS or not compact(values['length']):
                    skipped_rows.append({'row': rownum, 'reason': '마지막 합계 행'})
                    continue
            try:
                for key in ('material', 'spec'):
                    if values[key] is None:
                        raise ValueError('재질 또는 규격이 비어 있습니다. 병합 셀도 각 행에 값을 입력해 주세요.')
                    values[key] = str(values[key]).strip()
                values['length'] = millimeters(values['length'])
                raw_qty = values['quantity']
                if isinstance(raw_qty, bool):
                    raise ValueError('수량은 양의 정수여야 합니다.')
                qty = Decimal(re.sub(r'\s+', '', str(raw_qty)).replace(',', ''))
                if not qty.is_finite() or not 0 < qty <= 5000 or qty != qty.to_integral_value():
                    raise ValueError('수량은 1~5,000 사이의 정수여야 합니다.')
                values['quantity'] = int(qty)
                values['source_row'] = rownum
                values.update({name: re.sub(r'\s+', ' ', str(row[col])).strip() if row[col] is not None else ''
                               for name, col in metadata.items()})
                items.append(BOMItem(**values))
            except (ValueError, InvalidOperation, TypeError, ValidationError) as exc:
                message = str(exc)
                if isinstance(exc, ValidationError):
                    message = '; '.join(f"{e['loc'][0]}: {e['msg']}" for e in exc.errors())
                errors.append({'row': rownum, 'message': message[:250],
                               'title': str(row[metadata['title']] or '').strip() if 'title' in metadata else '',
                               'spec': str(values['spec'] or '').strip(), 'length': str(values['length'] or '')})
                if len(errors) >= 50:
                    break
        if errors:
            return {'items': [], 'errors': errors}
        if not items:
            raise ValueError('BOM 데이터가 없습니다. 시트와 헤더 행을 확인해 주세요.')
        if len(items) > 1000 or sum(i.quantity for i in items) > 5000:
            raise ValueError('BOM은 최대 1,000행, 총 5,000개 부재까지 지원합니다.')
        return {'items': [i.model_dump(mode='json') for i in items], 'errors': [], 'skipped_rows': skipped_rows}
    finally:
        wb.close()
