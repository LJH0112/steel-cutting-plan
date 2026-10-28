"""Print the shared report tables without changing the calculation engine."""
import io
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A3, A4, landscape
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from .reports import (MATERIAL_COLUMNS, COMBINATION_COLUMNS, MATERIAL_WIDTHS, COMBINATION_WIDTHS,
                      REPORT_ROW_HEIGHT, REPORT_HEADER_HEIGHT, REPORT_FONT_SIZE, REPORT_PAGE_MARGIN,
                      build_report, combination_rows)

ROOT = Path(__file__).resolve().parent.parent
FONT = 'NanumGothic'
pdfmetrics.registerFont(TTFont(FONT, str(ROOT / 'assets/fonts/NanumGothic-Regular.ttf')))


def safe_cell(value):
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def paper(value):
    if value not in ('A4', 'A3'):
        raise ValueError('인쇄 용지는 A4 또는 A3를 선택해 주세요.')
    return value


def excel_widths(ratios, paper_size, orientation):
    # Excel's default Calibri 11 column unit is seven pixels plus five pixels
    # padding. Preserve the same relative widths as the browser and PDF.
    size = A3 if paper_size == 'A3' else A4
    if orientation == 'landscape':
        size = landscape(size)
    # Match the PDF's printable width so fit-to-page does not shrink the fixed
    # row heights and font. Leave one point for rounding and the outer rules.
    pixels = (size[0] - 2 * REPORT_PAGE_MARGIN - 1) * 4 / 3
    return [round((pixels * v / sum(ratios) - 5) / 7, 3) for v in ratios]


def report_cell(cell, *, horizontal=None):
    cell.font = Font(name='맑은 고딕', size=REPORT_FONT_SIZE)
    cell.alignment = Alignment(horizontal=horizontal or ('right' if isinstance(cell.value, (float, int)) else 'left'),
                               vertical='center', wrap_text=False, shrink_to_fit=True)
    if isinstance(cell.value, (float, int)):
        cell.number_format = '#,##0.###'


def styled_sheet(wb, title, headers, rows, widths, paper_size='A4', orientation='portrait', report_title=None, metadata=None):
    ws = wb.create_sheet(title)
    if report_title:
        ws.page_margins = PageMargins(left=REPORT_PAGE_MARGIN/72, right=REPORT_PAGE_MARGIN/72,
                                      top=REPORT_PAGE_MARGIN/72, bottom=REPORT_PAGE_MARGIN/72,
                                      header=0, footer=14/72)
        meta = metadata or {}
        ws.append([report_title]); ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws.append(['[공 사 명] ' + meta.get('project_name', '')]); ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers)-2)
        ws.cell(2,len(headers)-1,'[날 짜] ' + meta.get('report_date','')); ws.merge_cells(start_row=2, start_column=len(headers)-1, end_row=2, end_column=len(headers))
        ws.append(['[공사구분] ' + meta.get('project_section', '')]); ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(headers))
    ws.append(headers)
    header_row = ws.max_row
    for row in rows:
        ws.append([safe_cell(v) for v in row])
    ws.freeze_panes = f'A{header_row+1}'
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = f'A{header_row}:{ws.cell(ws.max_row,len(headers)).coordinate}'
    ws.print_title_rows = f'1:{header_row}'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if orientation not in ('portrait','landscape'):
        raise ValueError('인쇄 방향은 portrait 또는 landscape여야 합니다.')
    ws.page_setup.orientation = orientation
    ws.page_setup.paperSize = ws.PAPERSIZE_A3 if paper(paper_size) == 'A3' else ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.left.text = '[파 일 명] ' + (metadata or {}).get('source_filename', '').replace('&','&&')
    ws.oddFooter.right.text = '&P / &N'
    ws.oddFooter.center.text = 'STEELPLAN'
    ws.oddFooter.center.size = 9
    for row in ws:
        for c in row:
            c.font = Font(name='맑은 고딕', size=10)
            c.alignment = Alignment(vertical='center', wrap_text=True)
            if isinstance(c.value, (float, int)):
                c.number_format = '#,##0.###'
    for cell in ws[header_row]:
        cell.font = Font(name='맑은 고딕', bold=False, color='000000')
        cell.fill = PatternFill('solid', fgColor='FFFF00')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(top=Side(style='thin'),bottom=Side(style='thin'),left=Side(style='thin'),right=Side(style='thin'))
    for row in ws.iter_rows(min_row=header_row+1):
        for c in row:
            c.border=Border(left=Side(style='hair'),right=Side(style='hair'))
    ws.row_dimensions[header_row].height=24
    if report_title:
        ws.cell(1,1).font=Font(name='맑은 고딕',size=20)
        ws.cell(1,1).alignment=Alignment(horizontal='center')
        ws.row_dimensions[1].height=38
        ws.row_dimensions[2].height=22
        ws.row_dimensions[3].height=22
        ws.row_dimensions[header_row].height = REPORT_HEADER_HEIGHT
        for row in ws.iter_rows(min_row=header_row + 1):
            ws.row_dimensions[row[0].row].height = REPORT_ROW_HEIGHT
            for cell in row:
                report_cell(cell)
        for cell in ws[header_row]:
            cell.font = Font(name='맑은 고딕', size=REPORT_FONT_SIZE)
    if not report_title: ws.row_dimensions[1].height = 26
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    return ws


def summary_text(plan, report):
    s = plan['summary']
    physical = report['output_piece_count']
    text = (f"발주 {s['bar_count']}본 | 총 로스율 {s['loss_rate']}% | "
            f"BOM 완성 부재 {s.get('input_piece_count', s['piece_count'])}개 = "
            f"계산 {s['piece_count']}개 + 제외 {s.get('excluded_piece_count', 0)}개 | 제작 절편 {physical}개")
    if plan.get('splices'):
        text += ' | ' + ('검토 필요 · 제작 확정 아님' if plan['splice_summary']['review_required'] else '이음 상세 사용자 승인 입력됨')
    return text


def notes_text(plan):
    basis = {'unspecified': '미지정(공란)', 'row_total': 'WEIGHT.KG 행 전체 중량', 'per_piece': 'WEIGHT.KG 개당 중량'}[plan['settings'].get('weight_basis', 'unspecified')]
    return (f"단위 mm, kg. 원자재의 조합길이·중량·잔철·사용율은 규격 소계에 표시; 소계·합계의 길이는 총 발주 길이. "
            f"사용율=제작 절편 길이/발주 길이×100. 총 로스율은 잔철+날 손실 기준. "
            f"원자재 중량은 단위중량 미제공으로 공란. 조합중량 해석: {basis}; 누락·오류가 있으면 소계도 공란. "
            f"조합길이·잔철은 1본 기준. 길이·등분은 같은 길이의 부재 행에 공통 적용; "
            f"같은 길이의 총수량 합계=등분×원자재 수량. 비고는 부재 번호만, 수량은 총수량 열에 표시.")


def export_xlsx(plan, paper_size='A4', orientation='portrait'):
    paper(paper_size)
    report = plan.get('report') or build_report(plan)
    wb = Workbook(); wb.remove(wb.active)
    meta = plan.get('report_metadata', {})
    ws = styled_sheet(wb, '원자재 리스트', [label for _, label in MATERIAL_COLUMNS],
                      [[row[key] for key, _ in MATERIAL_COLUMNS] for row in report['material_display_rows']],
                      excel_widths(MATERIAL_WIDTHS, paper_size, orientation), paper_size, orientation, '원자재 물량 집계표', meta)
    for index, row in enumerate(report['material_display_rows'], 5):
        if row['kind'] != 'detail':
            for cell in ws[index]:
                cell.font = Font(name='맑은 고딕', bold=True, size=REPORT_FONT_SIZE)
                cell.fill = PatternFill('solid', fgColor='FFFF00')
                cell.border = Border(top=Side(style='thin'), bottom=Side(style='thin'), left=Side(style='thin'), right=Side(style='thin'))
            report_cell(ws.cell(index, 2), horizontal='center')
    ws.append([summary_text(plan, report)])
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=10)
    ws.row_dimensions[ws.max_row].height = 32
    ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True)
    ws.append([notes_text(plan)])
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=10)
    ws.row_dimensions[ws.max_row].height = 45
    ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True)
    combo = styled_sheet(wb, '조합 리스트', [label for _, label in COMBINATION_COLUMNS], [],
                         excel_widths(COMBINATION_WIDTHS, paper_size, orientation), paper_size, orientation, '조합 리스트', meta)
    for group in report['combination_groups']:
        combo.append([f"[재질] {group['material']}   [규격] {group['spec']}   [단위중량]"])
        group_row = combo.max_row
        combo.merge_cells(start_row=group_row, start_column=1, end_row=group_row, end_column=11)
        combo.cell(group_row, 1).fill = PatternFill('solid', fgColor='FFFFFF')
        combo.cell(group_row, 1).font = Font(name='맑은 고딕', bold=True)
        combo.row_dimensions[group_row].height = REPORT_ROW_HEIGHT
        report_cell(combo.cell(group_row, 1))
        for pattern in group['patterns']:
            start = combo.max_row + 1
            printed = list(combination_rows(pattern))
            for row in printed:
                combo.append([safe_cell(row[key]) for key, _ in COMBINATION_COLUMNS])
            end = combo.max_row
            for col in range(1, 7):
                if end > start:
                    combo.merge_cells(start_row=start, start_column=col, end_row=end, end_column=col)
            for offset, row in enumerate(printed):
                if row['length_rowspan'] > 1:
                    for col in (7, 8):
                        combo.merge_cells(start_row=start+offset, start_column=col,
                                          end_row=start+offset+row['length_rowspan']-1, end_column=col)
            for row in combo.iter_rows(min_row=start, max_row=end):
                for cell in row:
                    report_cell(cell)
                    cell.border = Border(left=Side(style='thin'),right=Side(style='thin'))
                combo.row_dimensions[row[0].row].height = REPORT_ROW_HEIGHT
            for cell in combo[end]:
                cell.border=Border(left=Side(style='thin'),right=Side(style='thin'),bottom=Side(style='thin'))
        combo.append([safe_cell(group['total'][key]) for key,_ in COMBINATION_COLUMNS])
        combo.row_dimensions[combo.max_row].height = REPORT_ROW_HEIGHT
        for cell in combo[combo.max_row]:
            report_cell(cell)
            cell.fill=PatternFill('solid',fgColor='FFFF00')
            cell.border=Border(top=Side(style='thin'),bottom=Side(style='thin'),left=Side(style='thin'),right=Side(style='thin'))
    combo.auto_filter.ref = f'A4:K{combo.max_row}'
    styled_sheet(wb, 'BOM', ['재질', '규격', '길이 (mm)', '수량', '자재 종류', '원본 행', '부재 번호 (TITLE)',
                            '부재 분류 (ITEM)', 'WEIGHT.KG', 'PAINT.M', '계산 여부'],
                 [[i['material'], i['spec'], float(i['length']), i['quantity'], i.get('material_type', ''),
                   i.get('source_row'), i.get('title', ''), i.get('item_class', ''), i.get('weight', ''),
                   i.get('paint_area', ''), '포함' if i.get('included', True) else '제외'] for i in plan['bom']],
                 [18, 32, 18, 12, 18, 12, 24, 24, 18, 18, 15], paper_size)
    cfg = plan['settings']
    settings = [['요약', summary_text(plan, report)], ['설명', notes_text(plan)],
                ['절단날 손실 (mm/개)', cfg['kerf']], ['절단 후 잔재', '>30mm; 무절단 이음 통재는 절단 대상 아님'],
                ['계산 상태', '최적해 증명' if plan['summary']['optimal'] else '실행 가능 계획 / 최적해 미확정'],
                ['인쇄 용지', paper_size + (' 세로' if orientation=='portrait' else ' 가로')]]
    settings.extend([kind + ' 발주 길이', ', '.join(f'{v:g}' for v in values)] for kind, values in cfg.get('stock_lengths_by_type', {}).items())
    settings.extend([f"{r['material']} / {r['spec']}", r['detail']] for r in plan['statuses'])
    if plan.get('splices'):
        settings.extend([['이음 제작 여유 (mm/이음)', cfg['splice_allowance']],
                         ['재사용 잔재 최소 길이 (mm)', cfg['reusable_remnant_min']],
                         ['이음 여유 합계 (mm)', plan['summary']['total_splice_allowance']],
                         ['이음 여유 포함 순 로스율 (%)', plan['summary']['net_loss_rate']]])
    styled_sheet(wb, '계산 설정', ['항목', '값'], settings, [38, 130], paper_size)
    styled_sheet(wb, '계산 제외 부재', ['원본 행', '부재 번호', '자재 종류', '재질', '규격', '길이 (mm)', '수량', '허용 발주 길이 (mm)', '제외 원인'],
                 [[e['row'], e['title'], e['material_type'], e['material'], e['spec'], float(e['length']), e['quantity'],
                   ', '.join(f'{v:g}' for v in e['allowed_lengths']), e['display_reason']] for e in report['excluded']],
                 [12, 24, 18, 18, 32, 20, 12, 45, 65], paper_size)
    if plan.get('splices'):
        styled_sheet(wb, '이음 제작 검토', ['원본 행', '부재 번호', '재질', '규격', '완성 길이', '수량', '본재 길이',
                       '통재 본수/부재', '연결재 길이', '이음 수/부재', '여유/이음', '상태'],
                     [[a.get(k) for k in ('row', 'title', 'material', 'spec', 'finished_length', 'quantity', 'base_length',
                        'whole_bars_per_part', 'connector_length', 'joint_count', 'allowance_per_joint', 'status')]
                      for a in plan['splices']], [12, 22, 18, 32, 18, 12, 18, 18, 18, 18, 18, 22], paper_size)
        styled_sheet(wb, '잔재 재사용', ['원본 행', '부재 번호', '재질', '규격', '연결재 길이', '수량', '원자재 패턴', '사용 전 잔재', '사용 후 잔재'],
                     [[a.get(k) for k in ('row', 'title', 'material', 'spec', 'length', 'quantity', 'pattern_id', 'before', 'after')]
                      for a in plan['remnant_reuse']], [12, 22, 18, 32, 20, 12, 20, 20, 20], paper_size)
    styled_sheet(wb, '수량 검증', ['원본 행','부재 번호','재질','규격','BOM 수량','계획 완성 수량','제외 수량','제작 절편 수량','판정'],
        [[r[k] for k in ('row','title','material','spec','required_quantity','planned_quantity','excluded_quantity','component_quantity','status')]
         for r in report['reconciliation']], [12,24,18,32,16,18,16,18,22], paper_size, orientation)
    for sheet in wb:
        sheet.page_setup.orientation=orientation
    output = io.BytesIO(); wb.save(output)
    return output.getvalue()


def export_pdf(plan, paper_size='A4', orientation='portrait', section='all'):
    from .pdf_report import render_pdf
    paper(paper_size)
    report = plan.get('report') or build_report(plan)
    return render_pdf(plan, report, summary_text(plan, report), notes_text(plan), paper_size, orientation, section)


def sample_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = 'BOM'
    ws.append(['재질', '규격', '길이 (mm)', '수량'])
    for row in [('SS275', 'H-300×300×10×15', 3500, 6), ('SS275', 'H-300×300×10×15', 2400, 8),
                ('SS275', 'H-300×300×10×15', 1800, 6), ('SM355', 'H-400×200×8×13', 4200, 4),
                ('SM355', 'H-400×200×8×13', 2700, 6), ('SS275', 'C-200×80×7.5×11', 2900, 8),
                ('SS275', 'C-200×80×7.5×11', 1500, 10)]:
        ws.append(row)
    for col, width in [('A', 20), ('B', 32), ('C', 20), ('D', 15)]:
        ws.column_dimensions[col].width = width
    ws.freeze_panes = 'A2'
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
