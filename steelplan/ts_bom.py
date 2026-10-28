"""Read only TS BOM LIST and preserve the original Part Mark and row demand.

Ass Mark is an assembly reference, not a replacement for Part Mark. Assembly
summary rows are not cut parts, and their quantity is never multiplied into a
detail row. The first actual BOM supplied has its header on row 5.
"""
import hashlib
import re
from decimal import Decimal, InvalidOperation

from openpyxl.utils import get_column_letter
from pydantic import ValidationError

from .materials import classify_material, TS_BOM_MATERIAL_TYPES
from .models import BOMItem, ImportAudit, millimeters

SHEET = 'TS BOM LIST'
COLUMN_NAMES = {'material': 'grade', 'spec': 'profile', 'length': 'lengthmm',
                'quantity': 'qty', 'title': 'partmark', 'assembly_mark': 'assmark',
                'weight': 'weightkg', 'paint_area': 'aream'}


def normalized(value):
    # Handle m² in Area(m²), and Q'ty without confusing Ass Mark and Part Mark.
    return re.sub(r'[^a-z0-9]', '', str(value or '').lower())


def header_info(wb):
    ws = wb[SHEET]
    for index, row in enumerate(ws.iter_rows(max_row=30, max_col=100, values_only=True), 1):
        names = [normalized(c) for c in row]
        columns = {key: names.index(label) for key, label in COLUMN_NAMES.items() if names.count(label) == 1}
        if {'material', 'spec', 'length', 'quantity', 'title'} <= columns.keys():
            return {'name': SHEET, 'header_row': index,
                    'headers': [str(c).strip() if c is not None else '' for c in row],
                    'mapping': {k: columns[k] for k in ('material', 'spec', 'length', 'quantity')},
                    'source_columns': columns, 'locked': True, 'source_format': 'ts_bom'}
    raise ValueError('TS BOM LIST에서 Grade / PROFILE / Length(mm) / Q\'ty / Part Mark 헤더를 고유하게 찾지 못했습니다.')


def inspect_ts_bom(wb):
    return {'sheets': [header_info(wb)], 'suggested_sheet': SHEET,
            'source_format': 'ts_bom', 'excluded_sheets': [n for n in wb.sheetnames if n != SHEET]}


def locked_headers(wb, sheet, header_row):
    info = header_info(wb)
    if sheet != SHEET or header_row != info['header_row']:
        raise ValueError('TS BOM LIST 시트와 실제 헤더 행만 사용할 수 있습니다. 다른 시트·헤더로 변경할 수 없습니다.')
    return info


def parse_ts_bom(wb, data, sheet, header_row, mapping):
    info = locked_headers(wb, sheet, header_row)
    if mapping != info['mapping']:
        raise ValueError('TS BOM LIST의 PROFILE·Grade·Length(mm)·Q\'ty 열 매핑은 변경할 수 없습니다.')
    ws, columns = wb[SHEET], info['source_columns']
    if ws.max_row and ws.max_row > header_row + 10000:
        raise ValueError('TS BOM LIST 사용 범위가 10,000행을 초과합니다. 일부 행만 계산하지 않습니다.')
    excluded = [{'sheet': name, 'code': 'SHEET_NOT_SELECTED',
                 'reason': 'TS BOM LIST 이외 시트: 셀 데이터를 읽거나 최적화에 사용하지 않음'}
                for name in wb.sheetnames if name != SHEET]
    items, errors = [], []
    for rownum, row in enumerate(ws.iter_rows(min_row=header_row+1, max_col=100, values_only=True), header_row+1):
        if not any(str(v).strip() for v in row if v is not None):
            continue
        raw = {k: '' if row[col] is None else str(row[col]) for k, col in columns.items()}
        values = {k: v.strip() for k, v in raw.items()}
        material, spec, title = (values[k] for k in ('material', 'spec', 'title'))
        ass = values.get('assembly_mark', '')
        code = reason = None
        if normalized(spec) == 'profile' and normalized(title) == 'partmark':
            code, reason = 'REPEATED_HEADER', '반복 머리글'
        elif not spec and not title and any(normalized(c) in ('stotal', 'subtotal', 'total', 'grandtotal') for c in row):
            code, reason = 'SUBTOTAL', '원본 소계/합계행: 부재 수량 중복 집계 방지'
        elif spec and not title and ass and not material:
            code, reason = 'ASSEMBLY_SUMMARY', '조립 요약행(Part Mark·Grade 없음): 절단 부재와 중복 집계 방지'
        elif spec and classify_material(spec) not in TS_BOM_MATERIAL_TYPES:
            code, reason = 'OUTSIDE_PROFILE_SCOPE', 'PROFILE 대상 외: H빔·잔넬·C형강·환봉·T바만 최적화'
        if code:
            excluded.append({'sheet': SHEET, 'row': rownum, 'assembly_mark': ass, 'title': title,
                             'material': material, 'spec': spec, 'length': values['length'],
                             'quantity': values['quantity'], 'code': code, 'reason': reason})
            continue
        try:
            if not title:
                raise ValueError('Part Mark가 없습니다. Ass Mark로 대체하거나 임의 번호를 만들지 않습니다.')
            if isinstance(row[columns['quantity']], bool):
                raise ValueError('Q\'ty는 논리값이 아닌 정수여야 합니다.')
            qty = Decimal(values['quantity'])
            if not qty.is_finite() or qty != qty.to_integral_value() or not 0 < qty <= 5000:
                raise ValueError('Q\'ty는 1~5,000 사이의 정수여야 합니다.')
            item = BOMItem(source_format='ts_bom', material=material, spec=spec,
                           title=title, length=millimeters(values['length']), quantity=int(qty),
                           assembly_mark=ass, source_sheet=SHEET, source_row=rownum,
                           weight=values.get('weight', ''), paint_area=values.get('paint_area', ''),
                           source_cells={k: raw[k] for k in ('material','spec','title','length','quantity','assembly_mark') if k in raw})
            items.append(item)
        except (ValueError, InvalidOperation, ValidationError) as exc:
            errors.append({'row': rownum, 'title': title, 'spec': spec, 'length': values['length'], 'message': str(exc)[:250]})
            if len(errors) >= 50:
                break
    if errors:
        return {'items': [], 'errors': errors}
    if len(items) > 1000 or sum(i.quantity for i in items) > 5000:
        raise ValueError('대상 부재는 최대 1,000행, 총 5,000개까지 지원합니다. 일부 수량만 계산하지 않습니다.')
    audit = ImportAudit(header_row=header_row,
        columns={k: f'{get_column_letter(col+1)}: {info["headers"][col]}' for k, col in columns.items()},
        exclusions=excluded, source_sha256=hashlib.sha256(data).hexdigest())
    return {'items': [i.model_dump(mode='json') for i in items], 'errors': [], 'skipped_rows': [],
            'source_format': 'ts_bom', 'import_audit': audit.model_dump(mode='json')}
