import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app import RESULTS, app
from steelplan.bom import inspect_bom, parse_bom

client = TestClient(app)


def excel(rows, title='BOM'):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for row in rows:
        ws.append(row)
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def test_english_headers_offset_and_invalid_rows_not_silently_dropped():
    data = excel([['PROJECT BOM'], [], ['Grade', 'Section', 'Length (mm)', 'Qty'], ['SS275', 'H-300', 2000, 3],
                  ['SS275', 'H-300', 1500, 1.5]])
    s = inspect_bom(data)['sheets'][0]
    assert s['header_row'] == 3
    assert len(s['mapping']) == 4
    parsed = parse_bom(data, s['name'], s['header_row'], s['mapping'])
    assert parsed['items'] == []
    assert parsed['errors'][0]['row'] == 5


@pytest.mark.parametrize('qty', ['1e999999999', 0, -3, True, 'abc'])
def test_bad_quantities(qty):
    data = excel([['재질', '규격', '길이', '수량'], ['A', 'B', 2000, qty]])
    assert parse_bom(data, 'BOM', 1, {'material': 0, 'spec': 1, 'length': 2, 'quantity': 3})['errors']


def test_custom_headers_manual_mapping():
    data = excel([['meta'], ['mat_code', 'profile_code', 'cut_mm', 'amount'], ['A', 'B', '2,000.125', '2']])
    result = client.post('/api/headers', files={'file': ('bom.xlsx', data)}, data={'sheet': 'BOM', 'header_row': 2})
    assert result.status_code == 200
    parsed = parse_bom(data, 'BOM', 2, {'material': 0, 'spec': 1, 'length': 2, 'quantity': 3})
    assert parsed['items'][0]['length'] == '2000.125'
    assert parsed['items'][0]['quantity'] == 2
    with pytest.raises(ValueError):
        parse_bom(data, 'BOM', 2, {'material': 0, 'spec': 0, 'length': 2, 'quantity': 3})


def test_entire_workflow_sample_import_optimize_exports():
    assert client.get('/').status_code == 200
    assert client.get('/api/health').json() == {'status': 'ok'}
    assert client.get('/fonts/NanumGothic-Regular.ttf').status_code == 200
    sample = client.get('/api/sample').content
    inspected = client.post('/api/inspect', files={'file': ('bom.xlsx', sample)})
    assert inspected.status_code == 200
    s = inspected.json()['sheets'][0]
    response = client.post('/api/preview', files={'file': ('bom.xlsx', sample)},
                           data={'sheet': s['name'], 'header_row': s['header_row'], 'mapping': json.dumps(s['mapping'])})
    assert response.status_code == 200
    preview = response.json()
    assert preview['errors'] == []
    assert len(preview['items']) == 7
    response = client.post('/api/optimize', json={'items': preview['items'], 'kerf': 3, 'time_limit': 5})
    assert response.status_code == 200
    plan = response.json()
    assert plan['summary']['piece_count'] == 48
    assert plan['summary']['bar_count'] == 12
    assert all(p['remnant'] > 30 for p in plan['patterns'])
    xlsx = client.get(f"/api/results/{plan['job_id']}.xlsx")
    assert xlsx.status_code == 200
    wb = load_workbook(io.BytesIO(xlsx.content), data_only=False)
    assert wb.sheetnames == ['원자재 리스트', '조합 리스트', 'BOM', '계산 설정', '계산 제외 부재', '수량 검증', '발주 비교', '제작 발주 미확정', '승인 후보 기록', '보유 잔재 사용']
    assert next(row[3] for row in wb['원자재 리스트'].iter_rows(min_row=2, values_only=True) if row[1] == '합 계') == 12
    assert sum(row[3] for row in wb['BOM'].iter_rows(min_row=2, values_only=True)) == 48
    assert all(row[4] > 30 for row in wb['조합 리스트'].iter_rows(min_row=2, values_only=True) if isinstance(row[4], (float, int)))
    pdf = client.get(f"/api/results/{plan['job_id']}.pdf")
    assert pdf.status_code == 200
    assert pdf.content.startswith(b'%PDF-')
    assert b'/FontFile2' in pdf.content  # Korean TTF is embedded, not a reader font dependency.
    assert len(pdf.content) > 15000
    # Simulate expiry without sleeping.
    stamp, cached = RESULTS[plan['job_id']]
    RESULTS[plan['job_id']] = (stamp - 3601, cached)
    assert client.get(f"/api/results/{plan['job_id']}.xlsx").status_code == 404


def test_formula_injection_protection():
    response = client.post('/api/optimize', json={'items': [{'material': '=1+1', 'spec': 'H-300', 'length': 2000, 'quantity': 1}],
                                                 'kerf': 0, 'time_limit': 1})
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(client.get(f"/api/results/{response.json()['job_id']}.xlsx").content))
    assert wb['원자재 리스트']['A5'].data_type == 's'
    assert wb['원자재 리스트']['A5'].value.startswith("'")


def test_bad_files_and_bad_request():
    assert client.post('/api/inspect', files={'file': ('bom.xls', b'bad')}).status_code == 400
    assert client.post('/api/inspect', files={'file': ('bom.xlsx', b'bad')}).status_code == 400
    assert client.post('/api/optimize', json={'items': [], 'stock_lengths': [6000]}).status_code == 422
    result = client.post('/api/optimize', json={'items': [{'material': 'A', 'spec': 'B', 'length': 5970, 'quantity': 1}],
                                               'stock_lengths': [6000], 'kerf': 0})
    assert result.status_code == 400
    assert client.get('/api/results/not-found.pdf').status_code == 404
