import io
import json
from collections import Counter
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app import app
from steelplan import settings as settings_module
from steelplan.bom import inspect_bom, parse_bom
from steelplan.exports import export_pdf, export_xlsx
from steelplan.materials import DEFAULT_STOCK_LENGTHS, classify_material
from steelplan.models import OptimizeRequest, StockSettings
from steelplan.planning import PlanningIssues, assess_items, plan_by_material, validate_plan_input

client = TestClient(app)
HEADERS = ['ITEM', 'TITLE', 'MATL', 'DESCRIPTION', 'LENGTH', 'QTY', 'WEIGHT.KG', 'PAINT.M']


@pytest.fixture(autouse=True)
def isolated_settings(tmp_path, monkeypatch):
    monkeypatch.setattr(settings_module, 'SETTINGS_PATH', tmp_path / 'stock_settings.json')
    monkeypatch.delenv('STEELPLAN_ADMIN_PASSWORD', raising=False)


def actual_bom(rows, extra_sheet=False):
    wb = Workbook()
    ws = wb.active
    ws.title = 'TS ASSEM LIST'
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)
    if extra_sheet:
        other = wb.create_sheet('Other', 0)
        other.append(['재질', '규격', '길이', '수량'])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def item(spec, length=2000, quantity=1, **extra):
    return dict(material='SS275', spec=spec, length=length, quantity=quantity, **extra)


def settings():
    return StockSettings(stock_lengths_by_type=DEFAULT_STOCK_LENGTHS)


def confirmed(request, config):
    return request.model_copy(update={'exclude_invalid': True,
                                      'exclusion_token': validate_plan_input(request, config)['exclusion_token']})


@pytest.mark.parametrize(('spec', 'kind'), [(' H 300 * 300 * 10 * 15 ', 'H빔'), ('H-300×300×10×15', 'H빔'),
    ('P100*100*4', '각파이프'), ('P100*50*4', '각파이프'), ('[125*65*6*8', '잔넬'),
    ('L65*65*5', '앵글'), ('C150*75*20*3.2', 'C형강'), ('D19', '환봉'),
    ('PLT12*100', None), ('PIPE100', None), ('HUNKNOWN', None), ('D19ABC', None), ('P100*100', None)])
def test_explicit_material_classification(spec, kind):
    assert classify_material(spec) == kind


def test_exact_default_lengths():
    assert DEFAULT_STOCK_LENGTHS == {'H빔': list(range(8000, 16001, 1000)), '각파이프': [6000],
                                      '잔넬': [10000], '앵글': [10000], 'C형강': [10000], '환봉': [6000, 8000]}


def test_actual_bom_mapping_whitespace_metadata_and_total():
    data = actual_bom([[' COLUMN ', ' B-001 ', ' SS 275 ', ' H 300 * 300 * 10 * 15 ', ' 3500 ', ' 2 ', 10.5, 2.3],
                       [None] * 8, ['TOTAL', None, None, None, None, 2, 21, 4.6]], extra_sheet=True)
    inspected = inspect_bom(data)
    assert inspected['suggested_sheet'] == 'TS ASSEM LIST'
    s = next(s for s in inspected['sheets'] if s['name'] == 'TS ASSEM LIST')
    assert s['header_row'] == 1
    assert s['mapping'] == {'material': 2, 'spec': 3, 'length': 4, 'quantity': 5}
    parsed = parse_bom(data, s['name'], 1, s['mapping'])
    assert parsed['errors'] == []
    assert parsed['skipped_rows'] == [{'row': 4, 'reason': '마지막 합계 행'}]
    part = parsed['items'][0]
    assert part['source_row'] == 2
    assert part['title'] == 'B-001'
    assert part['item_class'] == 'COLUMN'
    assert part['material'] == 'SS275'
    assert part['spec'] == 'H300*300*10*15'
    assert part['weight'] == '10.5' and part['paint_area'] == '2.3'
    assert len(parsed['items']) == 1
    response = client.post('/api/preview', files={'file': ('ts.xlsx', data)}, data={
        'sheet': s['name'], 'header_row': 1, 'mapping': json.dumps(s['mapping'])})
    assert response.status_code == 200
    assert response.json()['issues'] == []
    assert response.json()['eligible_rows'] == 1


def test_real_part_named_total_is_not_omitted():
    data = actual_bom([['TOTAL', 'TOTAL', 'SS275', 'P100*100*4', 2000, 1]])
    s = inspect_bom(data)['sheets'][0]
    parsed = parse_bom(data, s['name'], 1, s['mapping'])
    assert len(parsed['items']) == 1 and parsed['skipped_rows'] == []


def test_numeric_cell_spaces_and_spaced_total_are_cleaned():
    data = actual_bom([['PIPE', 'P-01', 'SS275', 'P100*100*4', ' 2 900 ', ' 2 ', 10, 1],
                       [' TOTAL ', '', ' ', ' ', ' ', 2, 10, 1]])
    s = inspect_bom(data)['sheets'][0]
    parsed = parse_bom(data, s['name'], 1, s['mapping'])
    assert parsed['errors'] == []
    assert parsed['items'][0]['length'] == '2900'
    assert parsed['skipped_rows'][0]['row'] == 3


def test_rows_beyond_read_limit_are_not_silently_ignored():
    wb = Workbook(); ws = wb.active; ws.title = 'TS ASSEM LIST'
    ws.append(HEADERS)
    ws.append(['PIPE', 'P-01', 'SS275', 'P100*100*4', 2900, 1])
    ws.cell(12000, 1, 'DO NOT DROP')
    buffer = io.BytesIO(); wb.save(buffer)
    with pytest.raises(ValueError, match='10,000행'):
        parse_bom(buffer.getvalue(), ws.title, 1, {'material': 2, 'spec': 3, 'length': 4, 'quantity': 5})


def test_type_allowed_lengths_and_material_spec_isolation():
    parts = [item('H300*300*10*15', 7300, 3), item('P100*100*4', 2900, 4),
             item('[125*65*6*8', 4800, 2), item('L65*65*5', 4900, 2),
             item('C150*75*20*3.2', 4900, 2), item('D19', 3900, 2),
             {**item('P100*100*4', 2900, 1), 'material': 'SM355'}, item('P150*150*4', 2900, 1)]
    request = OptimizeRequest(items=parts, time_limit=5)
    plan = plan_by_material(request, settings())
    actual = Counter()
    for p in plan['patterns']:
        assert p['stock_length'] in DEFAULT_STOCK_LENGTHS[p['material_type']]
        assert Decimal(str(p['remnant'])) > 30
        product = Decimal(0)
        pieces = 0
        for c in p['cuts']:
            actual[p['material'], p['spec'], Decimal(str(c['length']))] += c['quantity'] * p['bar_count']
            product += Decimal(str(c['length'])) * c['quantity']
            pieces += c['quantity']
        assert product + request.kerf * pieces + Decimal(str(p['remnant'])) == Decimal(str(p['stock_length']))
    expected = Counter()
    for i in request.items:
        expected[i.material, i.spec, i.length] += i.quantity
    assert actual == expected
    assert plan['summary']['group_count'] == 8
    assert plan['excluded'] == []
    assert {o['material_type'] for o in plan['orders']} == set(DEFAULT_STOCK_LENGTHS)
    assert all(o['stock_length'] == 6000 for o in plan['orders'] if o['material_type'] == '각파이프')


def test_fixed_length_overflow_requires_explicit_exclusion_and_preserves_identity():
    parts = [item('P100*100*4', 6184, 2, source_row=7, title='P-OVER'),
             item('P100*100*4', 2900, 3, source_row=8, title='P-OK'),
             item('PLT12*100', 2000, 1, source_row=9, title='PL-01')]
    request = OptimizeRequest(items=parts, time_limit=3)
    with pytest.raises(PlanningIssues) as caught:
        plan_by_material(request, settings())
    exception = caught.value.issues[0]
    assert exception['row'] == 7 and exception['title'] == 'P-OVER'
    assert exception['spec'] == 'P100*100*4' and exception['length'] == '6184'
    assert exception['reason'] == '발주 가능 길이 초과'
    assert exception['allowed_lengths'] == [6000]
    response = client.post('/api/optimize', json=request.model_dump(mode='json', exclude_none=True))
    assert response.status_code == 409
    plan = plan_by_material(confirmed(request, settings()), settings())
    assert [e['title'] for e in plan['excluded']] == ['P-OVER', 'PL-01']
    assert plan['summary']['input_piece_count'] == 6
    assert plan['summary']['excluded_piece_count'] == 3
    assert plan['summary']['piece_count'] == 3
    assert sum(i['quantity'] for i in plan['bom']) == 6
    assert sum(i['quantity'] for i in plan['bom'] if i['included']) == 3
    assert all(p['stock_length'] == 6000 for p in plan['patterns'])
    assert not any(c['length'] == 6184 for p in plan['patterns'] for c in p['cuts'])
    wb = load_workbook(io.BytesIO(export_xlsx(plan)))
    assert wb['계산 제외 부재']['B2'].value == 'P-OVER'
    assert wb['계산 제외 부재']['F2'].value == 6184
    assert wb['계산 제외 부재']['I2'].value == '발주 길이 초과'
    assert wb['BOM']['K2'].value == '제외'
    assert export_pdf(plan).startswith(b'%PDF-')


def test_all_excluded_is_not_a_zero_test_success():
    with pytest.raises(ValueError, match='정상 부재가 없습니다'):
        plan_by_material(confirmed(OptimizeRequest(items=[item('P100*100*4', 6184)]), settings()), settings())


@pytest.mark.parametrize(('length', 'kerf', 'expected'), [(5970, 0, 'CUTTING_ALLOWANCE'), (5967, 3, 'CUTTING_ALLOWANCE'),
                                                       ('5966.999', 3, None), (6184, 0, 'STOCK_LENGTH_EXCEEDED')])
def test_fixed_stock_kerf_and_strict_remnant_boundary(length, kerf, expected):
    request = OptimizeRequest(items=[item('P100*100*4', length)], kerf=kerf)
    good, issues = assess_items(request.items, request.kerf, settings())
    assert (issues[0]['code'] if issues else None) == expected
    assert bool(good) == (expected is None)


def test_settings_admin_auth_atomic_save_and_new_lengths(monkeypatch):
    monkeypatch.setenv('STEELPLAN_ADMIN_PASSWORD', 'test-only-password')
    initial = client.get('/api/settings').json()
    initial['stock_lengths_by_type']['H빔'] = [8000]
    assert client.put('/api/settings', json=initial).status_code == 403
    assert client.put('/api/settings', json=initial, headers={'X-Admin-Password': 'wrong'}).status_code == 403
    result = client.put('/api/settings', json=initial, headers={'X-Admin-Password': 'test-only-password'})
    assert result.status_code == 200
    assert settings_module.load_settings().stock_lengths_by_type['H빔'] == [8000]
    assert settings_module.SETTINGS_PATH.exists()
    request = OptimizeRequest(items=[item('H-300', 6184)], time_limit=1)
    plan = plan_by_material(request, settings_module.load_settings())
    assert plan['orders'][0]['stock_length'] == 8000
    assert plan['excluded'] == []
    invalid = {**initial, 'stock_lengths_by_type': {**initial['stock_lengths_by_type'], 'H빔': [30]}}
    assert client.put('/api/settings', json=invalid, headers={'X-Admin-Password': 'test-only-password'}).status_code == 422
    assert settings_module.load_settings().stock_lengths_by_type['H빔'] == [8000]


def test_global_lengths_cannot_override_per_type():
    response = client.post('/api/optimize', json={'items': [item('P100*100*4', 6184)], 'stock_lengths': [12000]})
    assert response.status_code == 400
    assert '전역 발주 길이' in response.json()['detail']


def test_preview_overflow_not_dropped_and_unsupported_not_guessed():
    data = actual_bom([['PIPE', 'A-1', 'SS275', 'P100*100*4', 6184, 2],
                       ['PIPE', 'A-2', 'SS275', 'P100*100*4', 2900, 1],
                       ['PLATE', 'A-3', 'SS275', 'PLT12', 2000, 1], ['합계', None, None, None, None, 4]])
    s = inspect_bom(data)['sheets'][0]
    response = client.post('/api/preview', files={'file': ('ts.xlsx', data)}, data={
        'sheet': s['name'], 'header_row': 1, 'mapping': json.dumps(s['mapping'])})
    parsed = response.json()
    assert response.status_code == 200
    assert len(parsed['items']) == 3
    assert [e['row'] for e in parsed['issues']] == [2, 4]
    assert [e['code'] for e in parsed['issues']] == ['STOCK_LENGTH_EXCEEDED', 'UNSUPPORTED_MATERIAL']
    assert parsed['eligible_rows'] == 1
    result = client.post('/api/optimize', json={'items': parsed['items'], 'exclude_invalid': True,
                                              'exclusion_token': parsed['exclusion_token']})
    assert result.status_code == 200
    assert result.json()['summary']['excluded_piece_count'] == 3


def test_local_settings_without_password_and_cross_site_blocked():
    with TestClient(app, client=('127.0.0.1', 1234)) as local:
        value = settings().model_dump(mode='json')
        assert local.put('/api/settings', json=value).status_code == 200
        assert local.put('/api/settings', json=value, headers={'Sec-Fetch-Site': 'cross-site'}).status_code == 403


def test_stale_exclusion_consent_cannot_drop_newly_invalid_parts():
    request = OptimizeRequest(items=[item('P100*100*4', 6184), item('H300*300*10*15', 11000), item('D19', 2000)])
    approved = confirmed(request, settings())
    changed = {**DEFAULT_STOCK_LENGTHS, 'H빔': [8000]}
    with pytest.raises(PlanningIssues) as caught:
        plan_by_material(approved, StockSettings(stock_lengths_by_type=changed))
    assert len(caught.value.issues) == 2
    assert caught.value.validation['exclusion_token'] != approved.exclusion_token


def test_client_cannot_spoof_material_classification():
    response = client.post('/api/validate', json={'items': [item('PLT12', material_type='H빔')]})
    assert response.status_code == 200
    assert response.json()['issues'][0]['code'] == 'UNSUPPORTED_MATERIAL'
