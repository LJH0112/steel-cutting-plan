import io
import os
import random
import shutil
import subprocess
from collections import Counter
from decimal import Decimal
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import app
from steelplan.bom import inspect_bom, parse_bom
from steelplan.materials import DEFAULT_STOCK_LENGTHS
from steelplan.models import OptimizeRequest, StockSettings, RemnantStock
from steelplan.optimizer import optimize, units, MIN_REMAINDER
from steelplan.planning import plan_by_material, validate_plan_input
from steelplan.exports import export_xlsx,export_pdf


def part(length,qty=1,title='WBK3',**extra):
    return dict(material='SS275',spec='C100*50*20*2.3',assembly_mark='B5',title=title,length=length,quantity=qty,**extra)


def inventory(length=500,qty=1,confirmed=True,**extra):
    return dict(id='R1',material='SS275',spec='C100*50*20*2.3',length=length,quantity=qty,confirmed=confirmed,**extra)


def settings():return StockSettings(stock_lengths_by_type=DEFAULT_STOCK_LENGTHS)


def run(request):return plan_by_material(request,settings())


def approved(request,kind='SMALL_SPLICE'):
    v=validate_plan_input(request,settings())
    c=next(c for c in v['candidate_reviews'] if c['kind']==kind)
    return request.model_copy(update={'proposal_token':v['proposal_token'],'approved_candidate_ids':[c['id']]}),c


def assert_quantities(p):
    s=p['summary']
    assert s['input_piece_count']==s['piece_count']+s['pending_piece_count']
    assert s['total_stock_length']==pytest.approx(s['total_parts_length']+s['total_kerf_loss']+s['total_remnant'])
    assert s['total_parts_length']-s['total_splice_allowance']==pytest.approx(s['finished_parts_length'])
    assert sum(o['total_length'] for o in p['orders'])==s['total_new_stock_length']
    for r in p['report']['reconciliation']:
        assert r['required_quantity']==r['planned_quantity']+r['pending_quantity']+r['excluded_quantity']
    for pat in p['patterns']:
        cuts=sum(c['length']*c['quantity'] for c in pat['cuts'])
        assert cuts+pat['kerf_loss']+pat['remnant']==pytest.approx(pat['stock_length'])
        assert pat['remnant']>30 or pat.get('approval_id')


def test_new_length_priority_can_order_more_bars_to_buy_less():
    r=OptimizeRequest(items=[part(5960,2)],stock_lengths=[6000,13000],kerf=0,time_limit=2)
    p=optimize(r)
    assert (p['summary']['total_stock_length'],p['summary']['bar_count'])==(12000,2)
    old=optimize(r.model_copy(update={'objective':'bar_count'}))
    assert (old['summary']['total_stock_length'],old['summary']['bar_count'])==(13000,1)


def exhaustive(lengths,stocks,kerf):
    sizes=sorted([units(l)+units(kerf) for l in lengths],reverse=True)
    capacities=[units(s) for s in stocks];best=(10**20,10**20)
    def visit(i,bins):
        nonlocal best
        if i==len(sizes):
            value=(sum(next(s for s in capacities if s>=load+MIN_REMAINDER) for load in bins),len(bins))
            best=min(best,value);return
        seen=set()
        for n,load in enumerate(bins):
            if load not in seen and load+sizes[i]+MIN_REMAINDER<=capacities[-1]:
                seen.add(load);bins[n]+=sizes[i];visit(i+1,bins);bins[n]-=sizes[i]
        visit(i+1,bins+[sizes[i]])
    visit(0,[]);return best


def test_length_first_solver_against_independent_partition_oracle():
    rng=random.Random(742)
    for _ in range(20):
        lengths=[rng.randrange(500,5800,250) for _ in range(rng.randint(2,7))]
        stocks=[6000,9000,13000];kerf=rng.choice([0,3,4.5])
        r=OptimizeRequest(items=[part(l,title=str(i)) for i,l in enumerate(lengths)],stock_lengths=stocks,kerf=kerf,time_limit=2)
        p=optimize(r)
        assert p['summary']['optimal']
        assert (units(p['summary']['total_stock_length']),p['summary']['bar_count'])==exhaustive(lengths,stocks,kerf)


def test_finite_inventory_solver_against_independent_exhaustive_cost_oracle():
    rng=random.Random(310)
    for _ in range(15):
        lengths=[rng.choice([1000,2000,3000,4000]) for _ in range(rng.randint(2,6))]
        raw=rng.choice([3500,4500,6000]);stocks=[6000,9000]
        best=(10**20,10**20)
        def visit(i,bins,inventory_used):
            nonlocal best
            if i==len(lengths):
                fresh=[cap for cap,load,inv in bins if not inv]
                best=min(best,(sum(fresh),len(fresh)));return
            size=lengths[i]+3
            for j,(cap,load,inv) in enumerate(bins):
                if cap-load-size>30:
                    copy_bins=list(bins);copy_bins[j]=(cap,load+size,inv);visit(i+1,copy_bins,inventory_used)
            for cap in stocks:
                if cap-size>30:visit(i+1,bins+[(cap,size,False)],inventory_used)
            if not inventory_used and raw-size>30:visit(i+1,bins+[(raw,size,True)],True)
        visit(0,[],False)
        r=OptimizeRequest(items=[part(l,title=str(i)) for i,l in enumerate(lengths)],stock_lengths=stocks,time_limit=2)
        p=optimize(r,inventory=[RemnantStock(**inventory(raw))])
        assert p['summary']['optimal']
        assert (p['summary']['total_new_stock_length'],p['summary']['bar_count'])==best


def test_no_inventory_or_unverified_inventory_cannot_reduce_firm_orders():
    r=OptimizeRequest(items=[part(4900,2)],time_limit=1)
    p=run(r)
    unknown=run(r.model_copy(update={'remnant_inventory':[RemnantStock(**inventory(10000,confirmed=False))]}))
    assert p['summary']['total_new_stock_length']==unknown['summary']['total_new_stock_length']==10000
    assert unknown['inventory_usage']==[]
    assert unknown['comparison'][0]['total_new_stock_length']==unknown['comparison'][2]['total_new_stock_length']
    assert_quantities(unknown)


def test_verified_inventory_is_finite_and_jointly_packed_with_new_stock():
    r=OptimizeRequest(items=[part(4900,3)],remnant_inventory=[inventory(10000)],time_limit=2)
    p=run(r)
    assert p['summary']['piece_count']==3
    assert p['summary']['bar_count']==1
    assert p['summary']['total_new_stock_length']==10000
    assert p['inventory_usage'][0]['used_quantity']==1
    assert p['inventory_usage'][0]['used_length']==9806
    assert p['inventory_usage'][0]['remnant']==194
    assert_quantities(p)


@pytest.mark.parametrize(('change'),[{'material':'SM355'},{'spec':'C150*75*20*3.2'},{'confirmed':False}])
def test_wrong_material_spec_or_unverified_inventory_is_not_used(change):
    stock=inventory();stock.update(change)
    r=OptimizeRequest(items=[part(10050)],remnant_inventory=[stock],time_limit=1)
    v=validate_plan_input(r,settings());p=run(r)
    assert p['summary']['piece_count']==p['summary']['bar_count']==0
    assert p['summary']['pending_piece_count']==1
    c=next(c for c in v['candidate_reviews'] if c['kind']=='SMALL_SPLICE')
    assert not c['eligible_for_approval']
    with pytest.raises(ValueError):run(r.model_copy(update={'proposal_token':v['proposal_token'],'approved_candidate_ids':[c['id']]}))


def test_approved_short_splice_validates_allowance_kerf_and_reuses_the_same_bar():
    r=OptimizeRequest(items=[part(10050,2),part(2000,2,title='NORMAL')],remnant_inventory=[inventory(5000)],time_limit=2)
    before=run(r)
    assert before['summary']['pending_piece_count']==2
    assert before['summary']['piece_count']==2
    r,c=approved(r)
    assert c['shortage']==50 and c['connector_length']==55 and c['allowance']==5
    p=run(r)
    assert p['summary']['piece_count']==4 and p['summary']['pending_piece_count']==0
    assert p['summary']['bar_count']==2 and p['summary']['total_new_stock_length']==20000
    assert p['inventory_usage'][0]['used_quantity']==1
    assert p['summary']['total_splice_allowance']==10
    # 2 connectors and 2 normal parts share the very same finite 5m remnant.
    inv=next(pat for pat in p['patterns'] if pat.get('inventory_id'))
    assert {c['length']:c['quantity'] for c in inv['cuts']}=={2000:2,55:2}
    assert inv['kerf_loss']==12 and inv['remnant']==878
    assert_quantities(p)


@pytest.mark.parametrize(('shortage','expect_candidate'),[(100,True),(100.001,False)])
def test_100mm_boundary_and_large_shortages_are_never_silently_welded(shortage,expect_candidate):
    r=OptimizeRequest(items=[part(Decimal('10000')+Decimal(str(shortage)),2)],
        remnant_inventory=[inventory(2000)],splice_details_approved=True,time_limit=1)
    v=validate_plan_input(r,settings());p=run(r)
    assert bool([c for c in v['candidate_reviews'] if c['kind']=='SMALL_SPLICE'])==expect_candidate
    assert p['summary']['pending_piece_count']==2 and p['summary']['bar_count']==0
    assert p['pending'][0]['code']==('SMALL_SPLICE_REVIEW' if expect_candidate else 'LARGE_SPLICE_REVIEW')


def test_inventory_length_includes_saw_loss_allowance_and_strict_residual():
    for length,valid in [('88',False),('88.001',True)]:
        r=OptimizeRequest(items=[part(10050)],remnant_inventory=[inventory(length)],time_limit=1)
        # 55mm connector + 3mm kerf + >30mm remaining.
        c=validate_plan_input(r,settings())['candidate_reviews'][0]
        assert c['eligible_for_approval']==valid


def test_partial_inventory_capacity_keeps_remaining_demand_pending():
    r=OptimizeRequest(items=[part(10050,3)],remnant_inventory=[inventory(100)],time_limit=1)
    r,c=approved(r);assert c['quantity']==1
    p=run(r)
    assert p['summary']['piece_count']==1 and p['summary']['pending_piece_count']==2
    assert p['report']['reconciliation'][0]['required_quantity']==3
    assert_quantities(p)


def test_low_remnant_pattern_only_applies_after_that_pattern_is_approved():
    r=OptimizeRequest(items=[part(9990,10)],time_limit=1)
    before=run(r);assert before['summary']['piece_count']==0 and before['summary']['pending_piece_count']==10
    r,c=approved(r,'SHORT_REMNANT_PATTERN')
    assert c['remnant']==7 and c['bar_count']==10
    p=run(r)
    assert p['summary']['bar_count']==10 and p['summary']['piece_count']==10
    assert p['patterns'][0]['approval_id']==c['id']
    assert_quantities(p)


def test_multi_length_near_fit_is_offered_without_changing_unapproved_baseline():
    r=OptimizeRequest(items=[part(4982,2)],time_limit=1) # 9964 + 6 kerf = 9970 -> 30mm.
    before=run(r);r,c=approved(r,'SHORT_REMNANT_PATTERN');after=run(r)
    assert before['summary']['bar_count']==2 and after['summary']['bar_count']==1
    assert after['patterns'][0]['remnant']==30
    assert_quantities(after)


@pytest.mark.parametrize('change',[{'kerf':Decimal('4')},{'splice_allowance':Decimal('6')},{'transport_max_length':Decimal('9000')},
    {'remnant_inventory':[]}])
def test_stale_approvals_fail_when_cut_conditions_or_inventory_change(change):
    r=OptimizeRequest(items=[part(10050)],remnant_inventory=[inventory(500)],time_limit=1)
    r,_=approved(r)
    with pytest.raises(ValueError,match='변경'):
        run(r.model_copy(update=change))


def test_two_approvals_cannot_spend_the_same_inventory_twice():
    r=OptimizeRequest(items=[part(10050,title='A'),part(10050,title='B')],remnant_inventory=[inventory(100)],time_limit=1)
    v=validate_plan_input(r,settings())
    ids=[c['id'] for c in v['candidate_reviews'] if c['kind']=='SMALL_SPLICE']
    assert len(ids)==2
    with pytest.raises(ValueError,match='재고'):
        run(r.model_copy(update={'proposal_token':v['proposal_token'],'approved_candidate_ids':ids}))


def test_transport_and_processing_limits_keep_allowed_material_lengths():
    r=OptimizeRequest(items=[{**part(7300,3),'spec':'H300*300*10*15'}],transport_max_length=12000,processing_max_length=9000,time_limit=1)
    p=run(r)
    assert all(o['stock_length'] in [8000,9000] for o in p['orders'])
    assert all(o['stock_length'] in DEFAULT_STOCK_LENGTHS[o['material_type']] for o in p['orders'])
    assert_quantities(p)


def test_firm_orders_exports_pending_quantities_and_combined_marks(tmp_path):
    r=OptimizeRequest(items=[part(4900,2),part(9990,title='PENDING')],time_limit=1)
    p=run(r)
    assert {row['remarks'] for g in p['report']['combination_groups'] for pat in g['patterns'] for row in pat['rows']}=={'B5_WBK3'}
    wb=load_workbook(io.BytesIO(export_xlsx(p)))
    assert wb['제작 발주 미확정']['C2'].value=='PENDING'
    assert wb['제작 발주 미확정']['G2'].value==1
    assert wb['조합 리스트']['K6'].value=='B5_WBK3'
    assert wb['발주 비교']['C2'].value==10000
    assert sum(row[3] for row in wb['BOM'].iter_rows(min_row=2,values_only=True))==3
    path=tmp_path/'report.pdf';path.write_bytes(export_pdf(p))
    if shutil.which('pdftotext'):
        text=subprocess.check_output(['pdftotext','-raw',str(path),'-'],text=True)
        assert 'B5_WBK3' in text and 'PENDING' in text and '미확정' in text




def test_api_approves_only_the_displayed_snapshot_and_known_inventory():
    client=TestClient(app)
    request=OptimizeRequest(items=[part(10050)],remnant_inventory=[inventory(500)],time_limit=1).model_dump(mode='json')
    v=client.post('/api/validate',json=request)
    assert v.status_code==200
    v=v.json();c=next(c for c in v['candidate_reviews'] if c['kind']=='SMALL_SPLICE')
    args={**request,'proposal_token':v['proposal_token'],'approved_candidate_ids':[c['id']]}
    response=client.post('/api/optimize',json=args)
    assert response.status_code==200
    assert response.json()['summary']['piece_count']==1
    assert response.json()['manufacturing']['firm_orders_only']
    args['remnant_inventory'][0]['confirmed']=False
    assert client.post('/api/optimize',json=args).status_code==400


def test_all_pending_has_zero_firm_orders_but_exports_the_entire_bom():
    p=run(OptimizeRequest(items=[part(10895,5)],time_limit=1))
    assert p['orders']==[] and p['summary']['piece_count']==0
    assert p['summary']['pending_piece_count']==5
    assert sum(r['required_quantity'] for r in p['report']['reconciliation'])==5
    assert export_xlsx(p).startswith(b'PK') and export_pdf(p).startswith(b'%PDF-')


@pytest.mark.skipif(not os.environ.get('STEELPLAN_REAL_BOM'),reason='Set a local actual BOM path; customer data is not committed.')
def test_actual_ts_bom_all_292_parts_are_accounted_for_before_and_after_approval():
    data=Path(os.environ['STEELPLAN_REAL_BOM']).read_bytes();sheet=inspect_bom(data)['sheets'][0]
    parsed=parse_bom(data,sheet['name'],sheet['header_row'],sheet['mapping'])
    r=OptimizeRequest(items=parsed['items'],import_audit=parsed['import_audit'],time_limit=2)
    p=run(r)
    assert (p['summary']['input_piece_count'],p['summary']['piece_count'],p['summary']['pending_piece_count'])==(292,277,15)
    assert [(row['row'],row['quantity'],row['code']) for row in p['pending']]==[(459,10,'SHORT_REMNANT_REVIEW'),(464,5,'LARGE_SPLICE_REVIEW')]
    for row in p['report']['part_mark_reconciliation']:
        assert row['required_quantity']==row['planned_quantity']+row['pending_quantity']
    assert any(row['remarks']=='B5_WBK3' for g in p['report']['combination_groups'] for pp in g['patterns'] for row in pp['rows'])
    v=validate_plan_input(r,settings())
    c=next(c for c in v['candidate_reviews'] if c['kind']=='SHORT_REMNANT_PATTERN' and any(m['title']=='MPU6' for m in c['members']))
    p=run(r.model_copy(update={'proposal_token':v['proposal_token'],'approved_candidate_ids':[c['id']]}))
    assert (p['summary']['piece_count'],p['summary']['pending_piece_count'])==(287,5)
    assert p['pending'][0]['title']=='MPU7'
    assert_quantities(p)
    # Every direct planned source is counted by the actual row, including duplicate Part Marks.
    assert len(p['report']['reconciliation'])==141
    assert len(p['report']['part_mark_reconciliation'])==62
