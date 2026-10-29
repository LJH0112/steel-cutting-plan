import copy
import io
import re
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

from steelplan.exports import export_pdf, export_xlsx
from steelplan.materials import DEFAULT_STOCK_LENGTHS, classify_material
from steelplan.models import OptimizeRequest, StockSettings

from steelplan.planning import plan_by_material
from steelplan.reports import build_report


def legacy_request(**kwargs):
    """Preserve regression coverage for the explicitly requested compatibility mode."""
    kwargs.setdefault("planning_mode", "legacy")
    kwargs.setdefault("objective", "bar_count")
    return OptimizeRequest(**kwargs)




def make_plan(items, **kwargs):
    return plan_by_material(legacy_request(items=items, time_limit=2, **kwargs),
                            StockSettings(stock_lengths_by_type=DEFAULT_STOCK_LENGTHS))


def item(length, qty=1, title='B-01', **more):
    return {'material': 'SS275', 'spec': 'P100*100*4', 'length': length, 'quantity': qty, 'title': title, **more}


def test_long_pipe_defaults_to_splice_and_preserves_original_counts():
    p = make_plan([item(6184, 2, source_row=4)])
    assert p['excluded'] == []
    assert p['splices'][0]['connector_length'] == 189
    assert p['splices'][0]['review_required'] is True
    assert p['summary']['bar_count'] == 3  # 2 whole base bars + 1 connector stock.
    row = p['report']['reconciliation'][0]
    assert (row['required_quantity'], row['planned_quantity'], row['excluded_quantity'], row['component_quantity']) == (2, 2, 0, 4)
    assert row['row'] == 4 and row['title'] == 'B-01'


def test_duplicate_titles_and_lengths_are_audited_per_original_row():
    p = make_plan([item(2000, 2, 'SAME', source_row=2), item(2000, 3, 'SAME', source_row=3),
                   item(6184, 2, 'SAME', source_row=4)])
    rows = p['report']['reconciliation']
    assert [(r['row'], r['planned_quantity'], r['component_quantity'])for r in rows] == [(2,2,2),(3,3,3),(4,2,4)]
    assert sum(r['required_quantity'] for r in rows) == sum(r['planned_quantity'] + r['excluded_quantity'] for r in rows) == 7


def test_row_validation_catches_error_even_when_grand_total_matches():
    p = make_plan([item(2000, source_row=2), item(2000, source_row=3)])
    bad = copy.deepcopy(p)
    bad['bom'][1]['input_index'] = 0
    with pytest.raises(ValueError, match='원본 2행'):
        build_report(bad)


@pytest.mark.parametrize(('spec','expected'), [('TS100x100x3.2','각파이프'),('TS100*50*4','각파이프'),
    ('ROD19','환봉'),('ROD19.1','환봉'),('TSbad',None),('RODunknown',None)])
def test_reference_pdf_section_aliases_are_explicit(spec, expected):
    assert classify_material(spec) == expected


def test_reference_headers_yellow_subtotals_and_metadata_in_excel():
    p = make_plan([item(2900,2)], project_name='검증 공사', project_section='철골', source_filename='input.xlsx', report_date='2026-10-27')
    report = p['report']
    assert len(report['material_columns']) == 10 and report['material_columns'][-1]['label']=='비고'
    assert len(report['combination_columns']) == 11 and report['combination_columns'][-2]['label']=='총중량,KG'
    assert all(report['material_display_rows'][0][k] is None for k in ('combined_length','combined_weight','remnant','utilization'))
    assert report['material_display_rows'][1]['spec'] == '소 계'
    assert report['material_display_rows'][2]['spec'] == '합 계'
    wb = load_workbook(io.BytesIO(export_xlsx(p)))
    ws=wb['원자재 리스트']
    assert ws['A1'].value=='원자재 물량 집계표'
    assert '검증 공사' in ws['A2'].value and '철골' in ws['A3'].value
    assert ws['A4'].fill.fgColor.rgb.endswith('FFFF00')
    assert ws['A6'].fill.fgColor.rgb.endswith('FFFF00')
    assert ws['B6'].value=='소 계'
    assert ws.page_setup.orientation=='portrait'
    assert '[파 일 명] input.xlsx' == ws.oddFooter.left.text
    audit=wb['수량 검증']
    assert audit['E2'].value==audit['F2'].value==2


def test_pdf_separate_sections_metadata_and_total_page_numbers(tmp_path):
    if not shutil.which('pdftotext'): pytest.skip('poppler not installed')
    p = make_plan([item(2900,2)], project_name='FORM TEST', project_section='STEEL', source_filename='bom.xlsx')
    for section, title in [('materials','원자재 물량 집계표'),('combinations','조합 리스트')]:
        path=tmp_path/f'{section}.pdf'
        path.write_bytes(export_pdf(p,section=section))
        text=subprocess.check_output(['pdftotext','-layout',str(path),'-'],text=True)
        assert title in text and 'FORM TEST' in text and 'STEEL' in text and 'bom.xlsx' in text
        assert re.search(r'1\s*/\s*1',text)


def test_windows_launcher_pins_python_and_pauses_on_every_failure():
    path=Path(__file__).parent.parent/'run.bat'
    script=path.read_text(encoding='utf-8')
    assert 'py -3.13 -m venv' in script
    assert 'py -3 -m venv' not in script
    assert 'sys.version_info[:2] == (3, 13)' in script
    assert 'STEELPLAN_VENV=.venv-py313' in script
    assert 'if errorlevel 1 goto :error' in script
    for label in ('missing_python','wrong_venv','error'):
        branch=script.split('\n:'+label+'\n',1)[1].split('\n:',1)[0]
        assert branch.index('pause')<branch.index('exit /b 1')
    assert 'rmdir' not in script.lower() and 'del ' not in script.lower()


def test_pattern_with_many_wrapped_bom_names_can_continue_on_next_page():
    p = make_plan([item(500+i, title=f'TAG-{i:02d}-'+'긴번호'*26) for i in range(8)])
    assert p['report']['output_piece_count']==8
    assert export_pdf(p).startswith(b'%PDF-')
