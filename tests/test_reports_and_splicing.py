import copy
import io
import json
import re
import shutil
import subprocess
from collections import Counter
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import app
from steelplan.bom import inspect_bom, parse_bom
from steelplan.exports import export_pdf, export_xlsx
from steelplan.materials import DEFAULT_STOCK_LENGTHS
from steelplan.models import OptimizeRequest, StockSettings
from steelplan.planning import PlanningIssues, plan_by_material, validate_plan_input
from steelplan.reports import MATERIAL_COLUMNS, COMBINATION_COLUMNS, build_report


def part(length, quantity=1, title='A', weight='', **extra):
    return dict(material='SS275', spec='P100*100*4', length=length, quantity=quantity, title=title, weight=weight, **extra)


def plan(items, **extra):
    return plan_by_material(OptimizeRequest(items=items, time_limit=2, **extra),
                            StockSettings(stock_lengths_by_type=DEFAULT_STOCK_LENGTHS))


def pattern_plan(**extra):
    return plan([part(3000, 2, 'A-01', '150'), part(3000, 1, 'B-01', '75'), part(2000, 3, 'C-01', '100')], **extra)


def test_exact_columns_subtotals_and_merged_excel_patterns():
    p = pattern_plan()
    report = p['report']
    assert [c['label'] for c in report['material_columns']] == ['재질', '규격', '길이', '수량', '총중량,KG', '조합길이', '조합중량,KG', '잔철', '사용율', '비고']
    assert [c['label'] for c in report['combination_columns']] == ['번호', '원자재', '수량', '조합길이', '잔철', '사용율', '길이', '등분', '총수량', '총중량,KG', '비고']
    detail, subtotal, total = report['material_rows']
    assert detail['quantity'] == subtotal['quantity'] == total['quantity'] == 3
    assert detail['length'] == 6000 and subtotal['length'] == total['length'] == 18000
    assert total['combined_length'] == 15000
    assert total['utilization'] == 83.33
    assert detail['total_weight'] is None and detail['combined_weight'] is None
    patterns = report['combination_groups'][0]['patterns']
    assert len(patterns) == 1 and patterns[0]['quantity'] == 3
    assert sum(c['total_quantity'] for c in patterns[0]['cuts']) == 6
    assert all(c['total_quantity'] == c['split_count'] * patterns[0]['quantity'] for c in patterns[0]['cuts'])
    allocations = Counter()
    for c in patterns[0]['cuts']:
        for a in c['allocations']: allocations[a['title']] += a['quantity']
    assert allocations == {'A-01': 2, 'B-01': 1, 'C-01': 3}
    wb = load_workbook(io.BytesIO(export_xlsx(p)))
    assert [c.value for c in wb['원자재 리스트'][4]] == [label for _, label in MATERIAL_COLUMNS]
    assert [c.value for c in wb['조합 리스트'][4]] == [label for _, label in COMBINATION_COLUMNS]
    assert any(str(r) == 'A6:A7' for r in wb['조합 리스트'].merged_cells.ranges)
    assert sum(r[8] or 0 for r in wb['조합 리스트'].iter_rows(min_row=5, values_only=True) if r[0] != '합 계') == 6
    assert wb['원자재 리스트']['E5'].value is None


@pytest.mark.parametrize(('basis', 'expected'), [('unspecified', None), ('row_total', 325), ('per_piece', 675)])
def test_weights_only_follow_explicit_bom_interpretation(basis, expected):
    p = pattern_plan(weight_basis=basis)
    assert p['report']['material_rows'][0]['combined_weight'] == expected
    assert p['report']['material_rows'][-1]['combined_weight'] == expected
    assert p['report']['material_rows'][-1]['total_weight'] is None


@pytest.mark.parametrize('weight', ['', 'unknown', '-1', 'NaN'])
def test_incomplete_weight_does_not_become_a_partial_total(weight):
    p = plan([part(3000, 1, 'A', '100'), part(2000, 1, 'B', weight)], weight_basis='row_total')
    assert p['report']['material_rows'][-1]['combined_weight'] is None


def test_report_refuses_missing_or_extra_pieces():
    p = pattern_plan()
    broken = copy.deepcopy(p)
    broken['patterns'][0]['cuts'][0]['quantity'] += 1
    with pytest.raises(ValueError, match='수량'):
        build_report(broken)


@pytest.mark.parametrize(('size', 'width', 'height'), [('A4', 595.28, 841.89), ('A3', 841.89, 1190.55)])
@pytest.mark.parametrize('orientation', ['portrait','landscape'])
def test_paper_sizes_print_titles_and_page_numbers(size, width, height, orientation):
    if orientation == "landscape": width,height=height,width
    p = pattern_plan()
    wb = load_workbook(io.BytesIO(export_xlsx(p, size, orientation)))
    assert all(ws.page_setup.orientation == orientation for ws in wb)
    assert wb['원자재 리스트'].print_title_rows == '$1:$4'
    assert wb['BOM'].print_title_rows == '$1:$1'
    assert all(str(ws.page_setup.paperSize) == ('8' if size == 'A3' else '9') for ws in wb)
    assert '&P' in wb['조합 리스트'].oddFooter.right.text
    pdf = export_pdf(p, size, orientation)
    media = re.search(rb'/MediaBox\s*\[([^\]]+)\]', pdf)
    assert media
    values = [float(x) for x in media.group(1).split()]
    assert abs(values[2] - width) < .1 and abs(values[3] - height) < .1
    assert b'/FontFile2' in pdf
    assert int(re.search(rb'/Count\s+(\d+)', pdf).group(1)) >= 2


def test_real_format_sample_bom_and_export_quantities_match():
    from pathlib import Path
    data = (Path(__file__).parent.parent / 'examples/ts_assem_sample.xlsx').read_bytes()
    info = inspect_bom(data)['sheets'][0]
    parsed = parse_bom(data, info['name'], info['header_row'], info['mapping'])
    req = OptimizeRequest(items=parsed['items'])
    settings = StockSettings(stock_lengths_by_type=DEFAULT_STOCK_LENGTHS)
    validation = validate_plan_input(req, settings)
    p = plan_by_material(req.model_copy(update={'exclude_invalid': True, 'exclusion_token': validation['exclusion_token']}), settings)
    assert p['summary']['input_piece_count'] == 17
    assert p['summary']['piece_count'] == 16 and p['summary']['excluded_piece_count'] == 1
    output = sum(c['total_quantity'] for g in p['report']['combination_groups'] for pat in g['patterns'] for c in pat['cuts'])
    assert output == 17
    wb = load_workbook(io.BytesIO(export_xlsx(p)))
    assert sum(row[8] or 0 for row in wb['조합 리스트'].iter_rows(min_row=5, values_only=True) if row[0] != '합 계') == output


def test_explicit_strict_mode_does_not_split_or_weld_long_pipe():
    with pytest.raises(PlanningIssues) as error:
        plan([part(6184, title='LONG', source_row=8), part(2000, title='NORMAL')], enable_splice_planning=False)
    assert error.value.issues[0]['code'] == 'STOCK_LENGTH_EXCEEDED'
    assert error.value.issues[0]['row'] == 8 and error.value.issues[0]['quantity'] == 1
    assert error.value.issues[0]['length'] == '6184'


def check_splice_quantities(p):
    s = p['summary']
    assert s['input_piece_count'] == s['piece_count'] + s['excluded_piece_count']
    assert p['report']['output_piece_count'] == s['cut_piece_count']
    assert sum(pat['bar_count'] for pat in p['patterns']) == s['bar_count']
    assert Decimal(str(s['total_parts_length'])) - Decimal(str(s['total_splice_allowance'])) == Decimal(str(s['finished_parts_length']))
    assert Decimal(str(s['total_stock_length'])) == Decimal(str(s['total_parts_length'])) + Decimal(str(s['total_kerf_loss'])) + Decimal(str(s['total_remnant']))
    for pat in p['patterns']:
        if pat.get('uncut'):
            assert pat['parts_length'] == 6000 and pat['remnant'] == pat['kerf_loss'] == 0
        else:
            assert pat['remnant'] > 30


def test_opt_in_splice_6184_uses_6000_base_and_189_connector():
    p = plan([part(6184, 2, 'J-01')], enable_splice_planning=True)
    assert p['splices'][0]['base_length'] == 6000
    assert p['splices'][0]['connector_length'] == 189
    assert p['splices'][0]['review_required']
    assert not p['summary']['optimal']
    assert p['splice_summary']['whole_stock_bars'] == 2
    assert p['splice_summary']['extra_connector_stock_bars'] == 1
    assert p['summary']['bar_count'] == 3
    assert p['summary']['total_splice_allowance'] == 10
    assert p['excluded'] == []
    notes = ' '.join(c['remarks'] for g in p['report']['combination_groups'] for pat in g['patterns'] for c in pat['cuts'])
    assert '본재 6000' in notes and '연결재 189' in notes and 'J-01' in notes
    check_splice_quantities(p)
    assert export_pdf(p).startswith(b'%PDF-')


def test_connector_uses_matching_remnant_before_new_stock():
    p = plan([part(5600, title='N-01'), part(6184, title='J-01')], enable_splice_planning=True)
    assert p['splice_summary']['reused_connectors'] == 1
    assert p['splice_summary']['extra_connector_stock_bars'] == 0
    assert p['summary']['bar_count'] == 2
    audit = p['remnant_reuse'][0]
    assert audit['before'] == 397 and audit['after'] == 205
    assert audit['length'] == 189 and audit['title'] == 'J-01'
    check_splice_quantities(p)


@pytest.mark.parametrize('change', [{'material': 'SM355'}, {'spec': 'P150*150*4'}])
def test_other_material_or_section_remnant_cannot_be_used(change):
    p = plan([{**part(5600, title='N-01'), **change}, part(6184, title='J-01')], enable_splice_planning=True)
    assert p['splice_summary']['reused_connectors'] == 0
    assert p['splice_summary']['extra_connector_stock_bars'] == 1
    check_splice_quantities(p)


@pytest.mark.parametrize('change', [{'reusable_remnant_min': 500}, {'kerf': 100}])
def test_reuse_threshold_and_kerf_are_respected(change):
    p = plan([part(5600), part(6184)], enable_splice_planning=True, **change)
    assert p['splice_summary']['reused_connectors'] == 0
    check_splice_quantities(p)


@pytest.mark.parametrize('length', [6001, 6184, 11995, 12000, 18000, 99999])
def test_multi_joint_and_whole_connector_lengths_reconcile(length):
    p = plan([part(length)], enable_splice_planning=True)
    check_splice_quantities(p)


def test_splice_confirmation_only_from_explicit_user_input():
    p = plan([part(6184)], enable_splice_planning=True, splice_details_approved=True)
    assert not p['splice_summary']['review_required']
    assert p['splices'][0]['status'] == '사용자 승인 입력됨'


def test_square_pipe_setting_stays_6000():
    with pytest.raises(ValueError, match='6000mm 고정'):
        StockSettings(stock_lengths_by_type={**DEFAULT_STOCK_LENGTHS, '각파이프': [8000]})


def test_pdf_text_contains_headers_source_ids_and_exceptions(tmp_path):
    if not shutil.which('pdftotext'): pytest.skip('poppler not installed')
    p = pattern_plan()
    pdf = tmp_path / 'output.pdf'; pdf.write_bytes(export_pdf(p))
    extracted = subprocess.check_output(['pdftotext', '-layout', str(pdf), '-'], text=True)
    assert '원자재 물량 집계표' in extracted and '조합 리스트' in extracted
    assert 'A-01' in extracted and 'B-01' in extracted and 'C-01' in extracted
    assert '사용율' in extracted and '등분' in extracted and '합계' in ''.join(extracted.split())


def test_long_pdf_repeats_section_and_headers_without_losing_bom_numbers(tmp_path):
    if not shutil.which('pdftotext'): pytest.skip('poppler not installed')
    p = plan([part(3000 if i < 50 else 2000, title=f'BOM-{i:03d}-' + '긴부재번호' * 10)
              for i in range(100)])
    path = tmp_path / 'long.pdf'; path.write_bytes(export_pdf(p))
    text = subprocess.check_output(['pdftotext', '-layout', str(path), '-'], text=True)
    compact = ''.join(text.split())
    assert compact.count('[재질]SS275[규격]P100*100*4[단위중량]') >= 2
    assert all(f'BOM-{i:03d}' in text for i in range(100))
    assert p['report']['output_piece_count'] == 100


def test_api_a3_exports_and_invalid_size_rejected():
    client = TestClient(app)
    response = client.post('/api/optimize', json={'items': [part(2000)]})
    assert response.status_code == 200
    job = response.json()['job_id']
    assert client.get(f'/api/results/{job}.pdf?paper_size=A3').status_code == 200
    assert client.get(f'/api/results/{job}.xlsx?paper_size=A3').status_code == 200
    assert client.get(f'/api/results/{job}.pdf?paper_size=LETTER').status_code == 422
