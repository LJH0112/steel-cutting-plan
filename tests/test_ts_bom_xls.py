"""BIFF XLS import and Part Mark traces, using a synthetic, non-customer file."""
import io
import json
import os
import shutil
import subprocess
from collections import Counter
from pathlib import Path

import pytest
import xlrd
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app import app
from steelplan import settings as settings_module
from steelplan.bom import inspect_bom, parse_bom, read_headers
from steelplan.exports import export_pdf, export_xlsx
from steelplan.materials import DEFAULT_STOCK_LENGTHS, classify_material
from steelplan.models import OptimizeRequest, StockSettings
from steelplan.planning import PlanningIssues, plan_by_material, validate_plan_input

SAMPLE = Path(__file__).parent / 'fixtures/ts_bom_sample.xls'


@pytest.fixture(autouse=True)
def isolated_settings(tmp_path, monkeypatch):
    monkeypatch.setattr(settings_module, 'SETTINGS_PATH', tmp_path / 'settings.json')
    monkeypatch.delenv('STEELPLAN_ADMIN_PASSWORD', raising=False)


def parsed_file(path=SAMPLE):
    data = path.read_bytes()
    sheet = inspect_bom(data)['sheets'][0]
    return parse_bom(data, sheet['name'], sheet['header_row'], sheet['mapping'])


def configured_plan(parsed):
    request = OptimizeRequest(items=parsed['items'], import_audit=parsed['import_audit'], time_limit=2)
    # This is an explicit test setting, never an application default.
    config = StockSettings(stock_lengths_by_type={**DEFAULT_STOCK_LENGTHS, 'T바': [8000]})
    return plan_by_material(request, config)


def test_xls_only_reads_ts_bom_and_locks_real_columns(monkeypatch):
    seen = []
    original = xlrd.book.Book.sheet_by_name

    def trace(book, name):
        seen.append(name)
        return original(book, name)

    monkeypatch.setattr(xlrd.book.Book, 'sheet_by_name', trace)
    data = SAMPLE.read_bytes()
    info = inspect_bom(data)
    assert [s['name'] for s in info['sheets']] == ['TS BOM LIST']
    assert info['excluded_sheets'] == ['TS ASSEM LIST']
    sheet = info['sheets'][0]
    assert sheet['header_row'] == 5 and sheet['locked']
    assert sheet['mapping'] == {'material': 6, 'spec': 2, 'length': 3, 'quantity': 4}
    assert sheet['source_columns']['title'] == 1 and sheet['source_columns']['assembly_mark'] == 0
    parsed = parse_bom(data, 'TS BOM LIST', 5, sheet['mapping'])
    assert parsed['errors'] == []
    assert set(seen) == {'TS BOM LIST'}
    with pytest.raises(ValueError, match='실제 헤더'):
        read_headers(data, 'TS ASSEM LIST', 5)
    with pytest.raises(ValueError, match='매핑은 변경'):
        parse_bom(data, 'TS BOM LIST', 5, {'material': 6, 'spec': 2, 'length': 4, 'quantity': 3})
    with pytest.raises(ValueError, match='실제 헤더'):
        parse_bom(data, 'TS ASSEM LIST', 5, sheet['mapping'])


def test_scope_summary_rows_and_original_identifiers_are_preserved():
    parsed = parsed_file()
    items = parsed['items']
    assert sum(i['quantity'] for i in items) == 13  # No multiplication by assembly quantity.
    assert {classify_material(i['spec']) for i in items} == {'H빔','잔넬','C형강','환봉','T바'}
    original = items[0]
    assert original['title'] == 'WBK3' and original['assembly_mark'] == 'B5'
    assert original['quantity'] == 2 and original['length'] == '397.0'
    assert original['source_cells']['title'] == ' WBK3    '
    assert original['source_cells']['spec'] == 'H200*100*5.5*8        '
    assert original['source_cells']['material'] == '  SS275   '
    assert 'B5-WBK3' in {i['title'] for i in items}  # A literal source ID stays literal.
    reasons = Counter(e['code'] for e in parsed['import_audit']['exclusions'])
    assert reasons == {'SHEET_NOT_SELECTED':1, 'ASSEMBLY_SUMMARY':1, 'OUTSIDE_PROFILE_SCOPE':4, 'SUBTOTAL':2}
    assert not any(i['title'] in ('P-01','L-01','PL-01') for i in items)


@pytest.mark.parametrize(('spec', 'expected'), [('T100*50*5','T바'),('T100*50*5*7','T바'),
    ('TSBM16*45',None),('TS100*100*3.2','각파이프'),('T14',None)])
def test_t_bar_profile_is_explicit_and_not_an_assembly_mark(spec, expected):
    assert classify_material(spec) == expected


def test_t_bar_has_no_assumed_length_and_old_settings_migrate_without_overwriting():
    assert DEFAULT_STOCK_LENGTHS['T바'] == []
    old = {k:v for k,v in DEFAULT_STOCK_LENGTHS.items() if k != 'T바'}
    old['H빔'] = [12000]
    config = StockSettings(stock_lengths_by_type=old)
    assert config.stock_lengths_by_type['T바'] == [] and config.stock_lengths_by_type['H빔'] == [12000]
    parsed = parsed_file()
    request = OptimizeRequest(items=parsed['items'], import_audit=parsed['import_audit'])
    validation = validate_plan_input(request, config)
    assert [(e['title'],e['code'],e['allowed_lengths']) for e in validation['issues']] == [('T-01','STOCK_LENGTH_UNCONFIGURED',[])]
    with pytest.raises(PlanningIssues):
        plan_by_material(request, config)
    p = configured_plan(parsed)
    assert all(o['stock_length'] == 8000 for o in p['orders'] if o['material_type'] == 'T바')


def test_ts_bom_policy_cannot_reintroduce_square_pipe_into_the_plan():
    request = OptimizeRequest(items=[dict(source_format='ts_bom', source_sheet='TS BOM LIST',
        material='SS275',spec='P100*100*4',title='P-01',length=6184,quantity=1)])
    v = validate_plan_input(request, StockSettings(stock_lengths_by_type=DEFAULT_STOCK_LENGTHS))
    assert v['issues'][0]['code'] == 'OUTSIDE_PROFILE_SCOPE'
    assert not v['splice_candidates']


def test_member_totals_exports_and_original_cells_match_after_xls_import(tmp_path):
    parsed = parsed_file()
    p = configured_plan(parsed)
    expected, actual = Counter(), Counter()
    for i in parsed['items']:
        expected[i['material'],i['spec'],i['title']] += i['quantity']
    for g in p['report']['combination_groups']:
        for pattern in g['patterns']:
            for r in pattern['rows']:
                actual[g['material'],g['spec'],r['remarks']] += r['total_quantity']
    assert actual == expected
    assert p['summary']['piece_count'] == 13
    assert all(r['required_quantity'] == r['planned_quantity'] for r in p['report']['part_mark_reconciliation'])
    wb = load_workbook(io.BytesIO(export_xlsx(p)))
    assert '가져오기 제외' in wb and '원본 셀 추적' in wb and 'Part Mark 수량 검증' in wb
    assert wb['원본 셀 추적']['D2'].value == ' WBK3    '
    assert wb['원본 셀 추적']['C2'].value == 'B5'
    exported = Counter()
    for row in wb['조합 리스트'].iter_rows(min_row=5,values_only=True):
        if isinstance(row[8],int) and row[0] != '합 계':
            exported[row[10]] += row[8]
    assert exported == Counter({k[2]:v for k,v in expected.items()})
    pdf = tmp_path / 'xls-report.pdf'; pdf.write_bytes(export_pdf(p))
    if shutil.which('pdftotext'):
        text = subprocess.check_output(['pdftotext','-raw',str(pdf),'-'],text=True)
        assert 'WBK3' in text and 'B5-WBK3' in text
        assert '가져오기 제외 내역' in text and 'Part Mark 수량 검증' in text


def ts_xlsx(rows, sheet='TS BOM LIST', duplicate_title=False):
    w = Workbook();s=w.active;s.title=sheet
    headers=['Ass Mark','Part Mark','PROFILE','Length(mm)',"Q'ty",'Unit','Grade','Area(m²)','Weight(Kg)','Total(Kg)']
    if duplicate_title: headers[5] = 'Part Mark'
    s.append(headers)
    for row in rows:s.append(row)
    out=io.BytesIO();w.save(out);return out.getvalue()


@pytest.mark.parametrize('change', [{'title':''},{'material':''},{'length':'bad'},{'quantity':1.5},{'quantity':True}])
def test_target_rows_with_bad_values_block_entire_import(change):
    values=dict(title='A',material='SS275',length=2000,quantity=1)
    values.update(change)
    data=ts_xlsx([['B1',values['title'],'H200*100*5.5*8',values['length'],values['quantity'],'',values['material'],'','','']])
    result=parse_bom(data,'TS BOM LIST',1,{'material':6,'spec':2,'length':3,'quantity':4})
    assert result['items']==[] and result['errors']


def test_internal_spaces_not_rewritten_and_no_duplicate_header_guessing():
    data=ts_xlsx([['B1','A  B','H 200*100*5.5*8',2000,1,'','SS 275','','','']])
    result=parse_bom(data,'TS BOM LIST',1,{'material':6,'spec':2,'length':3,'quantity':4})
    assert result['items'][0]['title']=='A  B'
    assert result['items'][0]['material']=='SS 275' and result['items'][0]['spec']=='H 200*100*5.5*8'
    with pytest.raises(ValueError, match='고유하게'):
        inspect_bom(ts_xlsx([],duplicate_title=True))


def test_all_out_of_scope_rows_still_have_visible_exclusions():
    client=TestClient(app)
    data=ts_xlsx([['P1','P-01','P100*100*4',6184,1,'','SS275','','','']])
    response=client.post('/api/preview',files={'file':('bom.xlsx',data)},data={
        'sheet':'TS BOM LIST','header_row':1,'mapping':json.dumps({'material':6,'spec':2,'length':3,'quantity':4})})
    assert response.status_code==200
    d=response.json()
    assert d['items']==[] and d['eligible_rows']==0
    assert d['import_audit']['exclusions'][0]['title']=='P-01'


def test_xls_missing_ts_bom_does_not_fall_back_to_another_sheet():
    # Replace both UTF-16LE occurrences of the equal-length sheet name in BIFF
    # metadata and content, without introducing another XLS writer dependency.
    data=SAMPLE.read_bytes().replace('TS BOM LIST'.encode('utf-16le'),'XX BOM LIST'.encode('utf-16le'))
    # xlwt stores this ASCII name as BIFF compressed Unicode.
    data=data.replace(b'TS BOM LIST',b'XX BOM LIST')
    with pytest.raises(ValueError, match='반드시'):
        inspect_bom(data)


@pytest.mark.skipif(not os.environ.get('STEELPLAN_REAL_BOM'), reason='Set STEELPLAN_REAL_BOM to a local customer XLS; never commit it.')
def test_supplied_actual_bom_end_to_end():
    path=Path(os.environ['STEELPLAN_REAL_BOM'])
    data=path.read_bytes()
    client=TestClient(app)
    info=client.post('/api/inspect',files={'file':('BOM.xls',data)}).json()
    assert [s['name'] for s in info['sheets']]==['TS BOM LIST']
    s=info['sheets'][0]
    response=client.post('/api/preview',files={'file':('BOM.xls',data)},data={
        'sheet':s['name'],'header_row':s['header_row'],'mapping':json.dumps(s['mapping'])})
    assert response.status_code==200
    parsed=response.json()
    assert len(parsed['items'])==141 and sum(i['quantity'] for i in parsed['items'])==292
    assert [(e['row'],e['quantity'],e['code']) for e in parsed['issues']]==[
        (459,10,'CUTTING_ALLOWANCE'),(464,5,'STOCK_LENGTH_EXCEEDED')]
    request={'items':parsed['items'],'import_audit':parsed['import_audit'],'time_limit':2}
    assert client.post('/api/optimize',json=request).status_code==409
    response=client.post('/api/optimize',json={**request,'exclude_invalid':True,'exclusion_token':parsed['exclusion_token']})
    assert response.status_code==200
    p=response.json()
    assert p['summary']['piece_count']==277 and p['summary']['excluded_piece_count']==15
    assert len(p['report']['part_mark_reconciliation'])==62
    assert all(r['required_quantity']==r['planned_quantity']+r['excluded_quantity'] for r in p['report']['part_mark_reconciliation'])
    # Independent source-cell oracle: never create Ass Mark + Part Mark IDs,
    # infer quantities from assembly summaries, or pull parts from another sheet.
    s=xlrd.open_workbook(file_contents=data,on_demand=True).sheet_by_name('TS BOM LIST')
    source={i+1:s.row_values(i) for i in range(5,s.nrows)}
    for item in parsed['items']:
        r=source[item['source_row']]
        assert (item['title'],item['spec'],item['material'],item['quantity']) == (r[1].strip(),r[2].strip(),r[6].strip(),int(r[4]))
    actual=Counter()
    for g in p['report']['combination_groups']:
        for pattern in g['patterns']:
            for row in pattern['rows']:
                actual[g['material'],g['spec'],row['remarks']]+=row['total_quantity']
    assert actual==Counter({(r['material'],r['spec'],r['title']):r['planned_quantity'] for r in p['report']['part_mark_reconciliation'] if r['planned_quantity']})
    for fmt in ('xlsx','pdf'):
        output=client.get(f"/api/results/{p['job_id']}.{fmt}")
        assert output.status_code==200
        if fmt=='xlsx':
            wb=load_workbook(io.BytesIO(output.content))
            assert sum(row[3] for row in wb['Part Mark 수량 검증'].iter_rows(min_row=2,values_only=True))==292
            assert sum(row[4] for row in wb['Part Mark 수량 검증'].iter_rows(min_row=2,values_only=True))==277
        else:
            assert output.content.startswith(b'%PDF-')
