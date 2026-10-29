"""Member-separated reports preserve the physical plan and every BOM allocation."""
import copy
import io
import shutil
import subprocess
from collections import Counter
from xml.etree import ElementTree

import pytest
from openpyxl import load_workbook

from steelplan.exports import export_pdf, export_xlsx
from steelplan.materials import DEFAULT_STOCK_LENGTHS
from steelplan.models import OptimizeRequest, StockSettings

from steelplan.planning import plan_by_material
from steelplan.reports import (build_report, MATERIAL_WIDTHS, COMBINATION_WIDTHS,
                               REPORT_ROW_HEIGHT, REPORT_HEADER_HEIGHT, REPORT_PAGE_MARGIN)


def legacy_request(**kwargs):
    """Preserve regression coverage for the explicitly requested compatibility mode."""
    kwargs.setdefault("planning_mode", "legacy")
    kwargs.setdefault("objective", "bar_count")
    return OptimizeRequest(**kwargs)




def part(title, qty, length=2000, **extra):
    return dict(material='SS275', spec='P100*100*4', length=length,
                quantity=qty, title=title, **extra)


def plan(items, **extra):
    return plan_by_material(legacy_request(items=items, time_limit=2, **extra),
                            StockSettings(stock_lengths_by_type=DEFAULT_STOCK_LENGTHS))


def rows(p):
    return [row for group in p['report']['combination_groups']
            for pattern in group['patterns'] for row in pattern['rows']]


def test_same_length_members_get_individual_totals_without_changing_pattern():
    p = plan([part('A-01', 5), part('B-01', 1)])
    original = copy.deepcopy({k: v for k, v in p.items() if k != 'report'})
    report = build_report(p)
    assert {k: v for k, v in p.items() if k != 'report'} == original
    assert report['quantity_verified'] and report['member_quantity_verified']
    pattern = report['combination_groups'][0]['patterns'][0]
    assert (pattern['quantity'], pattern['stock_length'], pattern['combined_length'],
            pattern['remnant'], pattern['utilization']) == (3, 6000, 4000, 1994, 66.67)
    assert [(r['remarks'], r['total_quantity']) for r in pattern['rows']] == [('A-01', 5), ('B-01', 1)]
    # The per-bar count stays 2 for the whole length, even though A's 5 pieces
    # cannot be evenly divided across the three identical bars.
    assert all(r['split_count'] == 2 for r in pattern['rows'])
    assert [r['length_rowspan'] for r in pattern['rows']] == [2, 0]
    assert sum(r['total_quantity'] for r in pattern['rows']) == 2 * pattern['quantity']


def test_duplicate_numbers_and_missing_titles_preserve_source_rows():
    p = plan([part('SAME', 2, source_row=2), part('SAME', 3, source_row=3),
              part('', 1, source_row=4)])
    assert [(r['remarks'], r['total_quantity']) for r in rows(p)] == [('SAME', 5), ('', 1)]
    allocations = Counter()
    for row in rows(p):
        for source in row['allocations']:
            allocations[source['row']] += source['quantity']
    assert allocations == {2: 2, 3: 3, 4: 1}
    assert [r['planned_quantity'] for r in p['report']['reconciliation']] == [2, 3, 1]


def test_same_number_at_different_lengths_keeps_separate_member_rows():
    p = plan([part('MEMBER', 2, 3000), part('MEMBER', 2, 2000)])
    assert {(r['remarks'], r['length'], r['total_quantity']) for r in rows(p)} == {
        ('MEMBER', 3000, 2), ('MEMBER', 2000, 2)}


def test_member_weights_are_allocated_without_repeating_the_cut_total():
    p = plan([part('A-01', 2, 3000, weight='150'), part('B-01', 1, 3000, weight='75'),
              part('C-01', 3, 2000, weight='100')], weight_basis='row_total')
    assert {r['remarks']: r['total_weight'] for r in rows(p)} == {'A-01': 150, 'B-01': 75, 'C-01': 100}
    assert sum(r['total_weight'] for r in rows(p)) == p['report']['material_rows'][-1]['combined_weight'] == 325


def test_display_quantity_corruption_is_rejected(monkeypatch):
    import steelplan.reports as reports
    p = plan([part('A-01', 2), part('B-01', 2)])
    original = reports.combination_rows

    def broken(pattern):
        for row in original(pattern):
            yield {**row, 'total_quantity': row['total_quantity'] + 1}

    monkeypatch.setattr(reports, 'combination_rows', broken)
    with pytest.raises(ValueError, match='부재 번호별 출력 수량'):
        build_report(p)


def test_splice_rows_only_contain_member_numbers_and_audit_components_separately():
    p = plan([part('NORMAL', 1, 5600), part('JOINT-A', 2, 6184), part('JOINT-B', 1, 6184)])
    counts = Counter()
    for row in rows(p):
        counts[row['remarks']] += row['total_quantity']
    assert counts == {'NORMAL': 1, 'JOINT-A': 4, 'JOINT-B': 2}
    assert {r['remarks'] for r in rows(p)} == {'NORMAL', 'JOINT-A', 'JOINT-B'}
    assert sum(r['planned_quantity'] for r in p['report']['reconciliation']) == 4
    assert sum(counts.values()) == p['summary']['cut_piece_count'] == 7
    assert p['splice_summary']['review_required']


def test_excel_rows_widths_alignment_and_merges_are_fixed_for_long_names():
    long_title = 'LONG-' + '긴부재번호' * 25
    p = plan([part('A-01', 5), part(long_title, 1)])
    wb = load_workbook(io.BytesIO(export_xlsx(p)))
    ws = wb['조합 리스트']
    assert [(ws.cell(r, 11).value, ws.cell(r, 9).value) for r in (6, 7)] == [('A-01', 5), (long_title, 1)]
    assert {'A6:A7', 'G6:G7', 'H6:H7'} <= {str(r) for r in ws.merged_cells.ranges}
    assert ws['C6'].value == 3 and ws['H6'].value == 2  # Unchanged pattern counts.
    for name, widths in [('원자재 리스트', MATERIAL_WIDTHS), ('조합 리스트', COMBINATION_WIDTHS)]:
        sheet = wb[name]
        assert sheet.row_dimensions[4].height == REPORT_HEADER_HEIGHT
        for index, width in enumerate(widths, 1):
            dimension = sheet.column_dimensions[sheet.cell(4, index).column_letter]
            actual_pixels = dimension.width * 7 + 5
            printable_pixels = (595.2755905511812 - 2 * REPORT_PAGE_MARGIN - 1) * 4 / 3
            assert actual_pixels == pytest.approx(printable_pixels * width / sum(widths), abs=.01)
        assert sheet.page_margins.left == pytest.approx(REPORT_PAGE_MARGIN / 72)
    for r in (6, 7, 8):  # Both member rows and the group total.
        assert ws.row_dimensions[r].height == REPORT_ROW_HEIGHT
        assert ws.cell(r, 9).alignment.horizontal == 'right'
        assert ws.cell(r, 11).alignment.horizontal == 'left'
        assert ws.cell(r, 11).alignment.shrinkToFit
        assert not ws.cell(r, 11).alignment.wrap_text
    assert ws['K7'].value == long_title  # Full identifier retained in the file.


@pytest.mark.parametrize('orientation', ['portrait', 'landscape'])
def test_pdf_member_rows_use_distinct_lines_and_fixed_height(tmp_path, orientation):
    if not shutil.which('pdftotext'):
        pytest.skip('poppler not installed')
    p = plan([part('MEMBER-A', 5), part('MEMBER-B', 1)])
    path = tmp_path / 'members.pdf'
    path.write_bytes(export_pdf(p, orientation=orientation, section='combinations'))
    bbox = subprocess.check_output(['pdftotext', '-bbox', str(path), '-'], text=True)
    tree = ElementTree.fromstring(bbox[bbox.index('<html'):])
    words = {w.text: w for w in tree.iter() if w.tag.endswith('word') and w.text in ('MEMBER-A', 'MEMBER-B')}
    assert set(words) == {'MEMBER-A', 'MEMBER-B'}
    assert float(words['MEMBER-B'].get('yMin')) - float(words['MEMBER-A'].get('yMin')) == pytest.approx(REPORT_ROW_HEIGHT)
    assert float(words['MEMBER-A'].get('xMin')) == pytest.approx(float(words['MEMBER-B'].get('xMin')))
    text = subprocess.check_output(['pdftotext', '-layout', str(path), '-'], text=True)
    assert '×' not in text


def test_pdf_long_member_id_is_not_lost_in_fixed_height_cells(tmp_path):
    if not shutil.which('pdftotext'):
        pytest.skip('poppler not installed')
    title = 'LONG-' + '긴부재번호' * 25
    path = tmp_path / 'long-member.pdf'
    path.write_bytes(export_pdf(plan([part(title, 2)]), section='combinations'))
    text = subprocess.check_output(['pdftotext', '-raw', str(path), '-'], text=True)
    assert title in ''.join(text.split())


def test_real_bom_member_rows_and_excel_quantities_match_each_source():
    from pathlib import Path
    from steelplan.bom import inspect_bom, parse_bom
    from steelplan.planning import validate_plan_input
    data = (Path(__file__).parent.parent / 'examples/ts_assem_sample.xlsx').read_bytes()
    info = inspect_bom(data)['sheets'][0]
    parsed = parse_bom(data, info['name'], info['header_row'], info['mapping'])
    request = legacy_request(items=parsed['items'], time_limit=2)
    settings = StockSettings(stock_lengths_by_type=DEFAULT_STOCK_LENGTHS)
    validation = validate_plan_input(request, settings)
    p = plan_by_material(request.model_copy(update={'exclude_invalid': True, 'exclusion_token': validation['exclusion_token']}), settings)
    actual = Counter()
    for row in rows(p):
        actual[row['remarks']] += row['total_quantity']
    expected = Counter()
    for row in p['report']['reconciliation']:
        expected[row['title']] += row['component_quantity']
        assert row['required_quantity'] == row['planned_quantity'] + row['excluded_quantity']
    assert +actual == +expected
    wb = load_workbook(io.BytesIO(export_xlsx(p)))
    exported = Counter()
    for row in wb['조합 리스트'].iter_rows(min_row=5, values_only=True):
        if isinstance(row[8], int) and row[0] != '합 계':
            exported[row[10]] += row[8]
    assert exported == actual
